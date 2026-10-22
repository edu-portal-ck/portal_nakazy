"""Реєстр наказів — веб-застосунок для локальної мережі (версія — у changelog.json).

Ролі:
  * адміністратор — усе, довідники, імпорт/експорт, резервна копія;
  * секретар      — веде основні накази;
  * кадровик      — веде кадрові накази;
  * переглядач    — бачить усі реєстри, нічого не змінює;
  * без входу     — лише перегляд основних наказів.
Секретар і кадровик бачать усі реєстри, а редагують лише свій вид.

Запуск без Docker:  pip install -r requirements.txt && python app.py
Дані (база + скани + резервні копії) лежать у DATA_DIR (за замовчуванням ./data).
"""
import csv
import datetime as dt
import io
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import uuid
from functools import wraps

from flask import (Flask, Response, abort, g, jsonify, request, send_file,
                   send_from_directory, session)
from werkzeug.exceptions import HTTPException
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

BASE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.environ.get("DATA_DIR", os.path.join(BASE, "data"))
UPLOADS = os.path.join(DATA_DIR, "uploads")
BACKUPS = os.path.join(DATA_DIR, "backups")
DB_PATH = os.path.join(DATA_DIR, "orders.db")
INITIAL_ADMIN_PASSWORD = os.environ.get("INITIAL_ADMIN_PASSWORD", "admin")
PREVIEWS = os.path.join(DATA_DIR, "cache", "previews")  # PDF-версії Word-файлів (можна видаляти — створяться знову)
for _d in (UPLOADS, BACKUPS, PREVIEWS):
    os.makedirs(_d, exist_ok=True)

# Реєстри наказів — початкові значення (далі редагуються в програмі: Довідники → Реєстри).
# kind "main": основний реєстр (його бачать усі, веде секретар; поля: номер, дата, назва,
#              хто створив, хто підписав, статус, зв'язки, файли);
# kind "hr":   кадровий реєстр (бачать лише авторизовані, веде кадровик; ті самі поля).
# Зв'язки «скасовує / вносить зміни» — лише між наказами одного виду (кадрові номери не видно через основні).
# Поле «Примітка» прибрано з програми (колонку notes видаляє міграція m9).
# Кортеж: (ключ, вид, коротка назва для вкладки, повна назва, суфікс номера)
REGISTER_SEEDS = [
    ("main", "main", "Основна діяльність", "Реєстр наказів з основної діяльності", ""),
    ("hr1", "hr", "Кадрові", "Реєстр наказів з кадрових питань (особового складу)", "-к"),
]
STATUSES = {"active": "Чинний", "reserved": "Резерв", "void": "Втратив чинність"}
LINK_KINDS = {"cancels": "Скасовує", "amends": "Вносить зміни до"}
MIN_PASSWORD = 8
TRASH_DAYS = 30          # скільки днів наказ лежить у кошику до остаточного видалення
RESERVE_STALE_DAYS = 14  # резерв номера старший за це — у «Перевірці реєстру»

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 4 * 1024 ** 3  # 4 ГБ (скани та архіви для імпорту)


def _load_secret():
    path = os.path.join(DATA_DIR, "secret.key")
    if not os.path.exists(path):
        with open(path, "w") as f:
            f.write(uuid.uuid4().hex + uuid.uuid4().hex)
        try:
            os.chmod(path, 0o600)
        except OSError:
            pass
    with open(path) as f:
        return f.read().strip()


BEHIND_PROXY = os.environ.get("BEHIND_PROXY") == "1"   # програма за Caddy (HTTPS): див. docker-compose.https.yml

app.config.update(
    SECRET_KEY=_load_secret(),
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=BEHIND_PROXY,   # з HTTPS файл сеансу не передається відкритим каналом
    PERMANENT_SESSION_LIFETIME=dt.timedelta(hours=12),
)
if BEHIND_PROXY:
    from werkzeug.middleware.proxy_fix import ProxyFix
    app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)

# сторінка: лише власні скрипти/стилі (вбудовані в index.html), без чужих сайтів і без вбудовування в чужі сторінки
PAGE_CSP = ("default-src 'self'; script-src 'self' 'unsafe-inline'; style-src 'self' 'unsafe-inline'; "
            "img-src 'self' data: blob:; frame-src 'self' blob:; object-src 'none'; base-uri 'none'; "
            "form-action 'self'; frame-ancestors 'self'")


@app.before_request
def same_origin_only():
    """Зміни даних — лише із самої програми: запит із чужої сторінки (CSRF) відхиляється."""
    if request.method in ("GET", "HEAD", "OPTIONS") or not request.path.startswith("/api/"):
        return None
    from urllib.parse import urlsplit
    src = request.headers.get("Origin") or request.headers.get("Referer")
    if src and urlsplit(src).netloc.lower() != request.host.lower():
        return jsonify(error="Запит відхилено: він надійшов не зі сторінки програми"), 403
    return None


@app.after_request
def security_headers(resp):
    resp.headers.setdefault("X-Content-Type-Options", "nosniff")
    resp.headers.setdefault("X-Frame-Options", "SAMEORIGIN")
    resp.headers.setdefault("Referrer-Policy", "same-origin")
    # HSTS свідомо не вмикаємо: HTTPS можна вимкнути назад, а браузери пам'ятали б заборону HTTP рік
    return resp


# ---------------------------------------------------------------- Допоміжне
def now_iso():
    return dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def norm(s):
    return " ".join(str(s or "").split())


def split_names(s):
    return [n for n in (norm(p) for p in re.split(r"[;\n/]+", str(s or ""))) if n]


def upload_relpath(date, name):
    """Відносний шлях скану: РІК/МІСЯЦЬ/файл (за датою наказу)."""
    return f"{date[:4]}/{date[5:7]}/{name}"


def like_esc(v):
    """% і _ у пошуку шукаються як звичайні символи, а не як шаблон."""
    return v.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")


def seq_of(number):
    m = re.match(r"^\s*(\d+)", number or "")
    return int(m.group(1)) if m else None


def parse_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y-%m-%d")
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%y",
                "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M"):
        try:
            return dt.datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


# ---------------------------------------------------------------- БД
def connect():
    conn = sqlite3.connect(DB_PATH, timeout=15, isolation_level=None)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA foreign_keys=ON")
    # LOWER() у SQLite не працює з кирилицею — реєструємо власну функцію
    conn.create_function("ulower", 1, lambda v: v.casefold() if isinstance(v, str) else v)
    return conn


def get_db():
    if "db" not in g:
        g.db = connect()
    return g.db


@app.teardown_appcontext
def close_db(_exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def name_cache(db, table):
    return {norm(r["name"]).casefold(): r["id"] for r in db.execute(f"SELECT id, name FROM {table}")}


def ensure_name(db, table, cache, name):
    """Знаходить запис за ПІП або створює (для імпорту та міграції)."""
    name = norm(name)
    key = name.casefold()
    if key in cache:
        return cache[key]
    if table == "employees":
        cur = db.execute("INSERT INTO employees (name, active, created_at) VALUES (?,1,?)",
                         (name, now_iso()))
    else:
        cur = db.execute("INSERT INTO signers (name, active) VALUES (?,1)", (name,))
    cache[key] = cur.lastrowid
    return cur.lastrowid


# ---- Міграції. Нова зміна структури = нова функція m3, m4... і збільшення LATEST.
def m1_orders(conn):
    conn.execute("""
        CREATE TABLE orders (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            type TEXT NOT NULL,
            year INTEGER NOT NULL,
            seq INTEGER,
            number TEXT NOT NULL,
            date TEXT NOT NULL,
            title TEXT NOT NULL,
            responsible TEXT DEFAULT '',
            status TEXT NOT NULL DEFAULT 'active',
            notes TEXT DEFAULT '',
            file_name TEXT,
            file_stored TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )""")
    conn.execute("CREATE INDEX idx_orders_type_year ON orders(type, year)")
    conn.execute("CREATE INDEX idx_orders_date ON orders(date)")


def m2_roles_directories(conn):
    conn.execute("""
        CREATE TABLE employees (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            active INTEGER NOT NULL DEFAULT 1,
            is_admin INTEGER NOT NULL DEFAULT 0,
            is_secretary INTEGER NOT NULL DEFAULT 0,
            is_hr INTEGER NOT NULL DEFAULT 0,
            password_hash TEXT,
            must_change INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL
        )""")
    conn.execute("""
        CREATE TABLE signers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            position TEXT NOT NULL DEFAULT '',
            active INTEGER NOT NULL DEFAULT 1
        )""")
    conn.execute("""
        CREATE TABLE order_creators (
            order_id INTEGER NOT NULL REFERENCES orders(id) ON DELETE CASCADE,
            employee_id INTEGER NOT NULL REFERENCES employees(id),
            PRIMARY KEY (order_id, employee_id)
        )""")
    conn.execute("ALTER TABLE orders ADD COLUMN signer_id INTEGER REFERENCES signers(id)")
    conn.execute("ALTER TABLE orders ADD COLUMN created_by INTEGER")
    conn.execute("ALTER TABLE orders ADD COLUMN updated_by INTEGER")
    # тип «Інші» прибрано — такі накази стають основними
    conn.execute("UPDATE orders SET type='main' WHERE type NOT IN ('main','hr')")
    # старе текстове поле «Відповідальний» -> співробітники (довідник «хто створив»)
    cache = name_cache(conn, "employees")
    rows = conn.execute(
        "SELECT id, responsible FROM orders WHERE TRIM(COALESCE(responsible,'')) != ''").fetchall()
    for r in rows:
        for name in split_names(r["responsible"]):
            eid = ensure_name(conn, "employees", cache, name)
            conn.execute("INSERT OR IGNORE INTO order_creators (order_id, employee_id) VALUES (?,?)",
                         (r["id"], eid))


def m3_registers(conn):
    conn.execute("""
        CREATE TABLE registers (
            key TEXT PRIMARY KEY,
            kind TEXT NOT NULL,
            short TEXT NOT NULL,
            title TEXT NOT NULL,
            suffix TEXT NOT NULL DEFAULT '',
            sort INTEGER NOT NULL
        )""")
    for i, (key, kind, short, title, suffix) in enumerate(REGISTER_SEEDS):
        conn.execute("INSERT INTO registers (key, kind, short, title, suffix, sort) "
                     "VALUES (?,?,?,?,?,?)", (key, kind, short, title, suffix, i))
    # раніше кадрові накази були одним типом «hr» — переносимо в перший кадровий реєстр
    first_hr = next(k for k, kind, *_ in REGISTER_SEEDS if kind == "hr")
    conn.execute("UPDATE orders SET type=? WHERE type='hr'", (first_hr,))


def m4_reserve_and_folders(conn):
    # статус «Скасований» замінено на «Резерв»
    conn.execute("UPDATE orders SET status='reserved' WHERE status='cancelled'")
    # скани з плоскої папки -> РІК/МІСЯЦЬ (за датою наказу); якщо файл не вдалося перенести, лишається як є
    rows = conn.execute("SELECT id, date, file_stored FROM orders "
                        "WHERE file_stored IS NOT NULL AND file_stored NOT LIKE '%/%'").fetchall()
    for r in rows:
        rel = upload_relpath(r["date"], r["file_stored"])
        src, dst = os.path.join(UPLOADS, r["file_stored"]), os.path.join(UPLOADS, rel)
        try:
            if not os.path.exists(src):
                continue
            os.makedirs(os.path.dirname(dst), exist_ok=True)
            os.replace(src, dst)
        except OSError:
            continue
        conn.execute("UPDATE orders SET file_stored=? WHERE id=?", (rel, r["id"]))


def m5_many_files(conn):
    # до наказу можна приєднати кілька файлів (PDF, Word, Excel…): окрема таблиця
    conn.execute("""
        CREATE TABLE order_files (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            order_id INTEGER NOT NULL REFERENCES orders(id) ON DELETE CASCADE,
            name TEXT NOT NULL,
            stored TEXT NOT NULL,
            created_at TEXT NOT NULL
        )""")
    conn.execute("CREATE INDEX idx_order_files_order ON order_files(order_id)")
    conn.execute("INSERT INTO order_files (order_id, name, stored, created_at) "
                 "SELECT id, COALESCE(file_name, 'file'), file_stored, updated_at FROM orders "
                 "WHERE file_stored IS NOT NULL")
    conn.execute("UPDATE orders SET file_name=NULL, file_stored=NULL")


def m6_viewer_role(conn):
    # роль «Переглядач»: бачить усі реєстри, але нічого не змінює
    conn.execute("ALTER TABLE employees ADD COLUMN is_viewer INTEGER NOT NULL DEFAULT 0")


def m7_audit_trash_links(conn):
    # кошик: видалений наказ лишається в базі (з файлами) і остаточно зникає через TRASH_DAYS днів
    conn.execute("ALTER TABLE orders ADD COLUMN deleted_at TEXT")
    conn.execute("ALTER TABLE orders ADD COLUMN deleted_by INTEGER")
    # журнал дій: без зовнішніх ключів, щоб записи лишались і після остаточного видалення
    conn.execute("""
        CREATE TABLE audit (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            ts TEXT NOT NULL,
            user_id INTEGER,
            user_name TEXT NOT NULL DEFAULT '',
            action TEXT NOT NULL,
            order_id INTEGER,
            reg TEXT,
            number TEXT,
            year INTEGER,
            details TEXT NOT NULL DEFAULT ''
        )""")
    conn.execute("CREATE INDEX idx_audit_order ON audit(order_id)")
    # зв'язки: наказ order_id скасовує (cancels) або змінює (amends) наказ target_id
    conn.execute("""
        CREATE TABLE order_links (
            order_id INTEGER NOT NULL REFERENCES orders(id) ON DELETE CASCADE,
            target_id INTEGER NOT NULL REFERENCES orders(id) ON DELETE CASCADE,
            kind TEXT NOT NULL,
            PRIMARY KEY (order_id, target_id)
        )""")
    conn.execute("CREATE INDEX idx_links_target ON order_links(target_id)")
    # які нагадування «Перевірка реєстру» користувач уже бачив (тиждень / кінець місяця)
    conn.execute("ALTER TABLE employees ADD COLUMN checks_seen TEXT NOT NULL DEFAULT ''")


def m8_settings(conn):
    # налаштування програми, які змінює адміністратор (наприклад, доступ без входу)
    conn.execute("CREATE TABLE settings (key TEXT PRIMARY KEY, value TEXT NOT NULL)")


def m9_drop_notes(conn):
    # поле «Примітка» більше не використовується
    conn.execute("ALTER TABLE orders DROP COLUMN notes")


MIGRATIONS = {1: m1_orders, 2: m2_roles_directories, 3: m3_registers, 4: m4_reserve_and_folders,
              5: m5_many_files, 6: m6_viewer_role, 7: m7_audit_trash_links, 8: m8_settings,
              9: m9_drop_notes}
LATEST = max(MIGRATIONS)


def migrate():
    conn = connect()
    ver = conn.execute("PRAGMA user_version").fetchone()[0]
    has_orders = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name='orders'").fetchone()
    if ver == 0 and has_orders:  # база з першої версії програми (без номера версії)
        ver = 1
        conn.execute("PRAGMA user_version=1")
    if ver < LATEST:
        if has_orders:  # страхова копія бази перед зміною структури
            stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
            dst = sqlite3.connect(os.path.join(BACKUPS, f"orders_before_v{LATEST}_{stamp}.db"))
            conn.backup(dst)
            dst.close()
        for n in range(ver + 1, LATEST + 1):
            conn.execute("BEGIN IMMEDIATE")
            try:
                MIGRATIONS[n](conn)
                conn.execute(f"PRAGMA user_version={n}")
                conn.execute("COMMIT")
            except Exception:
                conn.execute("ROLLBACK")
                raise
    ensure_admin(conn)
    conn.close()


def ensure_admin(conn):
    """Якщо в системі немає жодного адміністратора з паролем — створює початкового."""
    n = conn.execute("SELECT COUNT(*) FROM employees WHERE is_admin=1 AND active=1 "
                     "AND password_hash IS NOT NULL").fetchone()[0]
    if n:
        return
    h = generate_password_hash(INITIAL_ADMIN_PASSWORD)
    row = conn.execute("SELECT id FROM employees WHERE name='Адміністратор'").fetchone()
    if row:
        conn.execute("UPDATE employees SET is_admin=1, active=1, password_hash=?, must_change=1 "
                     "WHERE id=?", (h, row["id"]))
    else:
        conn.execute("INSERT INTO employees (name, active, is_admin, password_hash, must_change, "
                     "created_at) VALUES ('Адміністратор',1,1,?,1,?)", (h, now_iso()))


migrate()


# ---------------------------------------------------------------- Помилки (JSON для /api)
MESSAGES = {401: "Потрібен вхід у систему", 403: "Недостатньо прав для цієї дії",
            404: "Не знайдено", 413: "Файл завеликий (перевищено ліміт розміру завантаження)",
            415: "Перегляд цього формату недоступний"}


@app.errorhandler(HTTPException)
def http_error(e):
    if request.path.startswith("/api/"):
        return jsonify(error=MESSAGES.get(e.code, e.description)), e.code
    return e


# ---------------------------------------------------------------- Автентифікація та права
def current_user():
    if "user" in g:
        return g.user
    g.user = None
    uid = session.get("uid")
    if uid:
        row = get_db().execute(
            "SELECT * FROM employees WHERE id=? AND active=1 AND password_hash IS NOT NULL",
            (uid,)).fetchone()
        if (row and (row["is_admin"] or row["is_secretary"] or row["is_hr"] or row["is_viewer"])
                and session.get("ph") == pw_stamp(row["password_hash"])):
            g.user = row
    return g.user


def pw_stamp(password_hash):
    """Відбиток пароля в сеансі: після зміни пароля всі інші сеанси цього користувача перестають діяти."""
    import hashlib
    return hashlib.sha256((password_hash or "").encode()).hexdigest()[:16]


def require_login():
    u = current_user()
    if not u:
        abort(401)
    return u


def require_admin():
    u = require_login()
    if not u["is_admin"]:
        abort(403)
    return u


def regs():
    """Реєстри наказів у порядку відображення: {ключ: {...}} (кеш на час запиту)."""
    if "regs" not in g:
        g.regs = {r["key"]: dict(r) for r in
                  get_db().execute("SELECT * FROM registers ORDER BY sort, key")}
    return g.regs


# Доступ без входу: full — основний реєстр зі сканами; list — накази повністю, але без сканів (скани після входу); none — нічого
PUBLIC_ACCESS = {"full": "Повний: основний реєстр і скани",
                 "list": "Без сканів: накази видно повністю, скани — після входу",
                 "none": "Закрито: без входу нічого не видно"}


def setting(key, default=""):
    if "settings" not in g:
        g.settings = {r["key"]: r["value"] for r in get_db().execute("SELECT key, value FROM settings")}
    return g.settings.get(key, default)


def public_access():
    v = setting("public_access", "full")
    return v if v in PUBLIC_ACCESS else "full"


def can_view(user, type_):
    r = regs().get(type_)
    if not r:
        return False
    if user is not None:
        return True
    return r["kind"] == "main" and public_access() != "none"


def can_open_files(user):
    return user is not None or public_access() == "full"


def can_edit(user, type_):
    r = regs().get(type_)
    if not user or not r:
        return False
    return bool(user["is_admin"] or (r["kind"] == "main" and user["is_secretary"])
                or (r["kind"] == "hr" and user["is_hr"]))


def user_json(u):
    if not u:
        return None
    return {
        "id": u["id"], "name": u["name"],
        "is_admin": bool(u["is_admin"]), "is_secretary": bool(u["is_secretary"]),
        "is_hr": bool(u["is_hr"]), "is_viewer": bool(u["is_viewer"]), "must_change": bool(u["must_change"]),
        "editable_types": [k for k in regs() if can_edit(u, k)],
    }


def audit(db, action, order=None, details="", user=None):
    """Запис у журнал дій. order — рядок/словник наказу (id, type, number, year) або None."""
    u = user if user is not None else current_user()
    if isinstance(details, (list, tuple)):
        details = "\n".join(x for x in details if x)
    db.execute("INSERT INTO audit (ts, user_id, user_name, action, order_id, reg, number, year, details) "
               "VALUES (?,?,?,?,?,?,?,?,?)",
               (now_iso(), u["id"] if u else None, u["name"] if u else "Система", action,
                order["id"] if order else None, order["type"] if order else None,
                order["number"] if order else None, order["year"] if order else None, details))


def short_text(s, n=120):
    s = norm(s)
    return s if len(s) <= n else s[:n - 1] + "…"


def fmt_d(iso):
    return f"{iso[8:10]}.{iso[5:7]}.{iso[:4]}" if iso and len(iso) >= 10 else (iso or "")


FAILS = {}  # employee_id -> (кількість помилок, час блокування до)


@app.get("/api/meta")
def meta():
    u = current_user()
    return jsonify(
        registers=[{k: r[k] for k in ("key", "kind", "short", "title", "suffix")}
                   for r in regs().values() if can_view(u, r["key"])],
        statuses=STATUSES,
        user=user_json(u),
        public_access=public_access(),
        files_open=can_open_files(u),
    )


@app.get("/api/login-list")
def login_list():
    rows = get_db().execute(
        "SELECT id, name FROM employees WHERE active=1 AND password_hash IS NOT NULL "
        "AND (is_admin=1 OR is_secretary=1 OR is_hr=1 OR is_viewer=1) ORDER BY name").fetchall()
    return jsonify(employees=[dict(r) for r in rows])


@app.post("/api/login")
def login():
    d = request.get_json(silent=True)
    d = d if isinstance(d, dict) else {}
    raw = d.get("employee_id")
    # id лише як ціле число: "1", "01", 1.0 — це той самий співробітник і той самий лічильник спроб
    if isinstance(raw, bool) or not isinstance(raw, (int, str)) or not re.fullmatch(r"\s*\d{1,9}\s*", str(raw)):
        return jsonify(error="Оберіть співробітника зі списку"), 400
    eid = int(str(raw).strip())
    pw = str(d.get("password") or "")
    if len(pw) > 200:
        return jsonify(error="Невірний пароль"), 401
    cnt, until = FAILS.get(eid, (0, 0))
    if until > time.time():
        mins = int((until - time.time()) / 60) + 1
        return jsonify(error=f"Забагато невдалих спроб. Спробуйте через {mins} хв."), 429
    row = get_db().execute("SELECT * FROM employees WHERE id=? AND active=1", (eid,)).fetchone()
    ok = (row and row["password_hash"] and (row["is_admin"] or row["is_secretary"] or row["is_hr"] or row["is_viewer"])
          and check_password_hash(row["password_hash"], pw))
    if not ok:
        cnt += 1
        FAILS[eid] = (0, time.time() + 120) if cnt >= 8 else (cnt, 0)
        if cnt >= 8 and row:
            audit(get_db(), "lock", details="8 невдалих спроб входу — вхід заблоковано на 2 хв",
                  user={"id": row["id"], "name": row["name"]})
        return jsonify(error="Невірний пароль"), 401
    FAILS.pop(eid, None)
    session.clear()
    session["uid"] = row["id"]
    session["ph"] = pw_stamp(row["password_hash"])
    session.permanent = True
    return jsonify(ok=True)


@app.post("/api/logout")
def logout():
    session.clear()
    return jsonify(ok=True)


@app.post("/api/password")
def change_password():
    u = require_login()
    d = request.get_json(silent=True) or {}
    if not check_password_hash(u["password_hash"], str(d.get("old") or "")):
        return jsonify(error="Поточний пароль невірний"), 400
    new = str(d.get("new") or "")
    if len(new) < MIN_PASSWORD:
        return jsonify(error=f"Пароль має містити щонайменше {MIN_PASSWORD} символів"), 400
    new_hash = generate_password_hash(new)
    get_db().execute("UPDATE employees SET password_hash=?, must_change=0 WHERE id=?",
                     (new_hash, u["id"]))
    audit(get_db(), "password", details="Змінено власний пароль")
    session["ph"] = pw_stamp(new_hash)   # поточний сеанс лишається, решта — ні
    return jsonify(ok=True)


# ---------------------------------------------------------------- Сторінка
@app.get("/")
def index():
    resp = send_from_directory(BASE, "index.html")
    resp.headers["Cache-Control"] = "no-cache"
    resp.headers["Content-Security-Policy"] = PAGE_CSP
    return resp


# ---------------------------------------------------------------- Накази
ORDER_SELECT = (
    "SELECT o.*, s.name AS signer_name, cb.name AS created_by_name, ub.name AS updated_by_name "
    "FROM orders o LEFT JOIN signers s ON s.id=o.signer_id "
    "LEFT JOIN employees cb ON cb.id=o.created_by "
    "LEFT JOIN employees ub ON ub.id=o.updated_by")


def creators_map(db):
    m = {}
    for r in db.execute("SELECT oc.order_id, e.id, e.name FROM order_creators oc "
                        "JOIN employees e ON e.id=oc.employee_id ORDER BY e.name"):
        m.setdefault(r["order_id"], []).append({"id": r["id"], "name": r["name"]})
    return m


def files_map(db):
    m = {}
    for r in db.execute("SELECT id, order_id, name, stored FROM order_files ORDER BY id"):
        m.setdefault(r["order_id"], []).append(
            {"id": r["id"], "name": r["name"], "ext": file_ext(r["name"], r["stored"]).lstrip(".")})
    return m


def links_maps(db):
    """Зв'язки між наказами (лише між ненакинутими в кошик): (вихідні, вхідні) за id наказу."""
    out, inc = {}, {}
    for r in db.execute(
            "SELECT l.order_id, l.target_id, l.kind, a.number AS a_num, a.date AS a_date, a.type AS a_type, "
            "b.number AS b_num, b.date AS b_date, b.type AS b_type FROM order_links l "
            "JOIN orders a ON a.id=l.order_id JOIN orders b ON b.id=l.target_id "
            "WHERE a.deleted_at IS NULL AND b.deleted_at IS NULL ORDER BY a.date, a.id"):
        out.setdefault(r["order_id"], []).append(
            {"id": r["target_id"], "kind": r["kind"], "number": r["b_num"], "date": r["b_date"], "type": r["b_type"]})
        inc.setdefault(r["target_id"], []).append(
            {"id": r["order_id"], "kind": r["kind"], "number": r["a_num"], "date": r["a_date"], "type": r["a_type"]})
    return out, inc


# наказ скасовано іншим (ненакинутим у кошик) наказом
CANCELLED_SQL = ("EXISTS (SELECT 1 FROM order_links l JOIN orders c ON c.id=l.order_id "
                 "WHERE l.target_id=o.id AND l.kind='cancels' AND c.deleted_at IS NULL)")


def serialize(row, cmap, user, fmap=None, lmaps=None):
    d = {k: row[k] for k in ("id", "type", "year", "seq", "number", "date", "title", "status",
                             "signer_id", "signer_name", "created_at", "updated_at")}
    d["creators"] = cmap.get(row["id"], [])
    d["files"] = (fmap or {}).get(row["id"], [])
    d["has_file"] = bool(d["files"])
    out, inc = lmaps or ({}, {})
    d["links"] = out.get(row["id"], [])
    d["linked_by"] = inc.get(row["id"], [])
    # «втратив чинність»: вручну або тому, що його скасував інший наказ
    d["void"] = row["status"] == "void" or any(x["kind"] == "cancels" for x in d["linked_by"])
    d["can_edit"] = can_edit(user, row["type"])
    if user and user["is_admin"]:   # «хто і коли вніс/змінив» бачить лише адміністратор
        d["created_by_name"] = row["created_by_name"]
        d["updated_by_name"] = row["updated_by_name"]
    return d


def fetch_order(db, oid, user):
    row = db.execute(ORDER_SELECT + " WHERE o.id=?", (oid,)).fetchone()
    return serialize(row, creators_map(db), user, files_map(db), links_maps(db)) if row else None


FIELD_NAMES = [("number", "Номер"), ("date", "Дата"), ("title", "Назва"), ("status", "Статус"),
               ("signer", "Хто підписав"), ("creators", "Хто створив"), ("links", "Зв'язки")]


def snapshot(db, oid):
    """Стан наказу людською мовою — для журналу дій."""
    o = db.execute(ORDER_SELECT + " WHERE o.id=?", (oid,)).fetchone()
    creators = [r["name"] for r in db.execute(
        "SELECT e.name FROM order_creators oc JOIN employees e ON e.id=oc.employee_id "
        "WHERE oc.order_id=? ORDER BY e.name", (oid,))]
    files = [r["name"] for r in db.execute("SELECT name FROM order_files WHERE order_id=? ORDER BY id", (oid,))]
    links = [f"{LINK_KINDS.get(r['kind'], r['kind'])} № {r['number']} від {fmt_d(r['date'])}" for r in db.execute(
        "SELECT l.kind, t.number, t.date FROM order_links l JOIN orders t ON t.id=l.target_id "
        "WHERE l.order_id=? ORDER BY t.date, t.id", (oid,))]
    return {"id": o["id"], "type": o["type"], "year": o["year"],
            "number": o["number"], "date": fmt_d(o["date"]), "title": o["title"] or "",
            "status": STATUSES.get(o["status"], o["status"]),
            "signer": o["signer_name"] or "", "creators": ", ".join(creators), "links": "; ".join(links),
            "files": files}


def describe_new(s):
    out = [f"{label}: {short_text(s[k])}" for k, label in FIELD_NAMES if s[k] and k not in ("status", "number")]
    if s["status"] != STATUSES["active"]:
        out.append(f"Статус: {s['status']}")
    if s["files"]:
        out.append("Файли: " + ", ".join(s["files"]))
    return out


def describe_changes(a, b):
    from collections import Counter
    out = [f"{label}: «{short_text(a[k]) or '—'}» → «{short_text(b[k]) or '—'}»"
           for k, label in FIELD_NAMES if a[k] != b[k]]
    added = list((Counter(b["files"]) - Counter(a["files"])).elements())
    removed = list((Counter(a["files"]) - Counter(b["files"])).elements())
    if added or removed:
        out.append("Файли: " + ", ".join([f"+ {x}" for x in added] + [f"− {x}" for x in removed]))
    return out


@app.get("/api/orders")
def list_orders():
    db = get_db()
    user = current_user()
    vis = [k for k in regs() if can_view(user, k)]
    where = [f"o.type IN ({','.join('?' * len(vis))})" if vis else "0", "o.deleted_at IS NULL"]
    params = list(vis)
    num = norm(request.args.get("num", "")).lstrip("№# ").strip()
    if num:
        if num.isdigit():
            # лише цифри = точний номер (1 знайде № 1 і 1-к, але не 11 чи 21)
            where.append("(ULOWER(o.number)=? OR ULOWER(o.number) LIKE ? ESCAPE '\\')")
            params += [num, num + "-%"]
        else:
            where.append("ULOWER(o.number) LIKE ? ESCAPE '\\'")
            params.append(f"%{like_esc(num.casefold())}%")
    q = norm(request.args.get("q", ""))
    if q:
        where.append("ULOWER(o.title) LIKE ? ESCAPE '\\'")
        params.append(f"%{like_esc(q.casefold())}%")
    t = request.args.get("type", "").strip()
    if t in vis:
        where.append("o.type=?")
        params.append(t)
    st = request.args.get("status", "").strip()
    if st == "void":
        where.append(f"(o.status='void' OR {CANCELLED_SQL})")
    elif st == "active":
        where.append(f"o.status='active' AND NOT {CANCELLED_SQL}")
    elif st in STATUSES:
        where.append("o.status=?")
        params.append(st)
    year = request.args.get("year", "").strip()
    if year.isdigit():
        where.append("o.year=?")
        params.append(int(year))
    month = request.args.get("month", "").strip()
    if month.isdigit() and 1 <= int(month) <= 12:
        where.append("substr(o.date, 6, 2)=?")
        params.append(f"{int(month):02d}")
    creator = request.args.get("creator", "").strip()
    if creator.isdigit():
        where.append("EXISTS (SELECT 1 FROM order_creators oc WHERE oc.order_id=o.id "
                     "AND oc.employee_id=?)")
        params.append(int(creator))
    signer = request.args.get("signer", "").strip()
    if signer.isdigit():
        where.append("o.signer_id=?")
        params.append(int(signer))
    sql = ORDER_SELECT
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY o.date DESC, COALESCE(o.seq, 0) DESC, o.id DESC"
    cmap = creators_map(db)
    fmap = files_map(db)
    lmaps = links_maps(db)
    rows = [serialize(r, cmap, user, fmap, lmaps) for r in db.execute(sql, params)]
    ykeys = [t] if t in vis else vis
    years = [r["year"] for r in db.execute(
        f"SELECT DISTINCT year FROM orders WHERE type IN ({','.join('?' * len(ykeys))}) "
        "AND deleted_at IS NULL ORDER BY year DESC", ykeys)] if ykeys else []
    return jsonify(orders=rows, years=years)


@app.get("/api/filter-options")
def filter_options():
    """Хто створював / підписував накази цього реєстру — для випадаючих фільтрів."""
    user = current_user()
    t = request.args.get("type", "")
    if not can_view(user, t):
        return jsonify(creators=[], signers=[])
    db = get_db()
    creators = db.execute(
        "SELECT DISTINCT e.id, e.name FROM order_creators oc JOIN orders o ON o.id=oc.order_id "
        "JOIN employees e ON e.id=oc.employee_id WHERE o.type=? AND o.deleted_at IS NULL "
        "ORDER BY e.name", (t,)).fetchall()
    signers = db.execute(
        "SELECT DISTINCT s.id, s.name FROM orders o JOIN signers s ON s.id=o.signer_id "
        "WHERE o.type=? AND o.deleted_at IS NULL ORDER BY s.name", (t,)).fetchall()
    return jsonify(creators=[dict(r) for r in creators], signers=[dict(r) for r in signers])


def next_number(db, type_, year):
    row = db.execute("SELECT MAX(seq) AS m FROM orders WHERE type=? AND year=? AND deleted_at IS NULL",
                     (type_, year)).fetchone()
    seq = (row["m"] or 0) + 1
    return seq, f"{seq}{regs()[type_]['suffix']}"


@app.get("/api/next-number")
def api_next_number():
    user = require_login()
    type_ = request.args.get("type", "")
    date = parse_date(request.args.get("date")) or dt.date.today().strftime("%Y-%m-%d")
    if type_ not in regs() or not can_edit(user, type_):
        return jsonify(number="")
    _, number = next_number(get_db(), type_, int(date[:4]))
    return jsonify(number=number)


def with_suffix(type_, number):
    """Номер лише з цифр отримує суфікс реєстру (1 -> 1-к у кадровому); інші номери — як вписано."""
    number = number.strip()
    return number + regs()[type_]["suffix"] if number.isdigit() else number


def duplicate_exists(db, type_, year, number, exclude_id=0):
    return db.execute(
        "SELECT id, number, date FROM orders WHERE type=? AND year=? AND ULOWER(number)=ULOWER(?) AND id!=? "
        "AND deleted_at IS NULL", (type_, year, number, exclude_id)).fetchone()


def int_list(values):
    out = []
    for v in values:
        try:
            i = int(v)
        except (TypeError, ValueError):
            continue
        if i not in out:
            out.append(i)
    return out


def read_form(db, user, fixed_type=None):
    f = request.form
    type_ = fixed_type or f.get("type", "").strip()
    date = parse_date(f.get("date"))
    title = f.get("title", "").strip()
    errors = []
    if type_ not in regs():
        errors.append("Оберіть реєстр")
    elif not can_edit(user, type_):
        errors.append("У вас немає прав вести накази цього реєстру")
    kind = regs().get(type_, {}).get("kind")
    status = f.get("status", "active")
    if status not in STATUSES:
        status = "active"
    if not date:
        errors.append("Вкажіть дату наказу")
    if not title and status != "reserved":
        errors.append("Вкажіть назву / зміст наказу")
    creators = int_list(f.getlist("creators"))
    if creators:
        found = {r["id"] for r in db.execute(
            f"SELECT id FROM employees WHERE id IN ({','.join('?' * len(creators))})", creators)}
        if len(found) != len(creators):
            errors.append("Невідомий співробітник у списку «Хто створив»")
    signer_id = None
    raw = f.get("signer_id", "").strip()
    if raw:
        try:
            signer_id = int(raw)
        except ValueError:
            signer_id = -1
        if not db.execute("SELECT 1 FROM signers WHERE id=?", (signer_id,)).fetchone():
            errors.append("Невідомий підписант")
    links = {}
    for v in f.getlist("links"):   # «kind:id»
        k, _, i = str(v).partition(":")
        if k not in LINK_KINDS or not i.isdigit():
            errors.append("Невірний зв'язок із наказом")
            continue
        t = db.execute("SELECT type FROM orders WHERE id=? AND deleted_at IS NULL", (int(i),)).fetchone()
        if not t or regs().get(t["type"], {}).get("kind") != kind:
            errors.append("Пов'язаний наказ не знайдено (можливо, його видалено) або він з реєстру іншого виду")
            continue
        links[int(i)] = k
    return {
        "type": type_, "kind": kind, "date": date, "title": title,
        "number": f.get("number", "").strip(),
        "status": status,
        "creators": creators, "signer_id": signer_id, "links": links,
    }, errors


def set_links(db, oid, links):
    db.execute("DELETE FROM order_links WHERE order_id=?", (oid,))
    for target, kind in links.items():
        if target != oid:
            db.execute("INSERT INTO order_links (order_id, target_id, kind) VALUES (?,?,?)", (oid, target, kind))


def stored_path(stored):
    """Повний шлях до скану; None, якщо шлях виходить за межі папки uploads."""
    path = os.path.normpath(os.path.join(UPLOADS, stored))
    return path if os.path.commonpath([path, os.path.normpath(UPLOADS)]) == os.path.normpath(UPLOADS) else None


WORD_EXT = {".doc", ".docx", ".odt", ".rtf", ".dot", ".dotx", ".docm", ".ott"}
CALC_EXT = {".xls", ".xlsx", ".xlsm", ".ods"}
PRES_EXT = {".ppt", ".pptx", ".pps", ".ppsx", ".odp"}
OFFICE_EXT = WORD_EXT | CALC_EXT | PRES_EXT
DIRECT_EXT = {".pdf", ".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".txt"}   # лише це відкривається в браузері
SAFE_MIME = {".pdf": "application/pdf", ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
             ".gif": "image/gif", ".webp": "image/webp", ".bmp": "image/bmp", ".txt": "text/plain"}
_convert_lock = threading.Lock()
_LO_PROFILE = os.path.join(tempfile.gettempdir(), "nakazy_lo_profile")


def file_ext(original, stored):
    return (os.path.splitext(original or "")[1] or os.path.splitext(stored or "")[1]).lower()


def preview_pdf(stored, ext):
    """PDF-версія Word-файлу: створюється через LibreOffice при першому перегляді й кешується.
    Повертає шлях до PDF або None, якщо перетворити не вдалося."""
    src = stored_path(stored)
    if not src or not os.path.exists(src):
        return None
    return convert_pdf(src, ext, os.path.join(PREVIEWS, os.path.basename(stored) + ".pdf"))


def convert_pdf(src, ext, out):
    """Перетворює офісний файл src на PDF out через LibreOffice (з кешем за часом зміни)."""
    def fresh():
        return os.path.exists(out) and os.path.getmtime(out) >= os.path.getmtime(src)

    if fresh():
        return out
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None
    with _convert_lock:  # LibreOffice не любить паралельних запусків із одним профілем
        if fresh():
            return out
        tmp = tempfile.mkdtemp(prefix="conv_", dir=PREVIEWS)
        try:
            work = os.path.join(tmp, "doc" + ext)  # проста назва без кирилиці
            shutil.copyfile(src, work)
            # для таблиць кожен аркуш — на одній сторінці, щоб широкі таблиці не розрізались
            target = ('pdf:calc_pdf_Export:{"SinglePageSheets":{"type":"boolean","value":"true"}}'
                      if ext in CALC_EXT else "pdf")
            subprocess.run([exe, "--headless", "--norestore", "--nolockcheck",
                            "-env:UserInstallation=file://" + _LO_PROFILE,
                            "--convert-to", target, "--outdir", tmp, work],
                           timeout=120, capture_output=True)
            pdf = os.path.join(tmp, "doc.pdf")
            if os.path.exists(pdf):
                os.replace(pdf, out)
                return out
        except (subprocess.TimeoutExpired, OSError):
            pass
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
    return None


def warm_preview(stored, original):
    """Готує PDF у фоні одразу після завантаження, щоб перший перегляд був швидким."""
    ext = file_ext(original, stored)
    if ext in OFFICE_EXT:
        threading.Thread(target=preview_pdf, args=(stored, ext), daemon=True).start()


def save_upload(file_storage, date):
    if not file_storage or not file_storage.filename:
        return None, None
    original = file_storage.filename
    base, ext = os.path.splitext(original)
    ext = re.sub(r"[^a-z0-9.]", "", ext.lower())[:10]  # secure_filename викидає кирилицю разом із розширенням
    stored = upload_relpath(date, f"{uuid.uuid4().hex}_{secure_filename(base) or 'file'}{ext}")
    path = stored_path(stored)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    file_storage.save(path)
    warm_preview(stored, original)
    return original, stored


def add_files(db, oid, file_list, date, saved=None):
    for f in file_list:
        name, stored = save_upload(f, date)
        if stored:
            if saved is not None:
                saved.append(stored)
            db.execute("INSERT INTO order_files (order_id, name, stored, created_at) VALUES (?,?,?,?)",
                       (oid, name, stored, now_iso()))


def remove_files(db, oid, ids=None, defer=None):
    """Видаляє файли наказу (усі, або лише з переліку id) з бази. З диска — одразу,
    а якщо передано список defer, то імена збираються туди й видаляються лише після успішного збереження."""
    for r in db.execute("SELECT id, stored FROM order_files WHERE order_id=?", (oid,)).fetchall():
        if ids is None or r["id"] in ids:
            if defer is None:
                remove_stored(r["stored"])
            else:
                defer.append(r["stored"])
            db.execute("DELETE FROM order_files WHERE id=?", (r["id"],))


def move_files(db, oid, date):
    """Змінилась дата наказу -> файли переїжджають у папку нового року/місяця."""
    for r in db.execute("SELECT id, stored FROM order_files WHERE order_id=?", (oid,)).fetchall():
        want = upload_relpath(date, os.path.basename(r["stored"]))
        src, dst = stored_path(r["stored"]), stored_path(want)
        if want == r["stored"] or not src or not dst or not os.path.exists(src):
            continue
        try:
            os.makedirs(os.path.dirname(dst), exist_ok=True)
            os.replace(src, dst)
        except OSError:
            continue
        for d in (os.path.dirname(src), os.path.dirname(os.path.dirname(src))):
            try:
                os.rmdir(d)
            except OSError:
                break
        db.execute("UPDATE order_files SET stored=? WHERE id=?", (want, r["id"]))


def remove_stored(stored):
    path = stored_path(stored) if stored else None
    if not path:
        return
    try:
        os.remove(path)
    except OSError:
        return
    import glob
    base = os.path.basename(stored)
    # PDF-перегляд і розпаковані з архіву файли (base__N.ext та їхні PDF)
    for p in [os.path.join(PREVIEWS, base + ".pdf")] + glob.glob(os.path.join(PREVIEWS, glob.escape(base) + "__*")):
        try:
            os.remove(p)
        except OSError:
            pass
    for d in (os.path.dirname(path), os.path.dirname(os.path.dirname(path))):  # прибрати порожні МІСЯЦЬ/РІК
        try:
            if os.path.normpath(d) != os.path.normpath(UPLOADS):
                os.rmdir(d)
        except OSError:
            break


def set_creators(db, oid, creators):
    db.execute("DELETE FROM order_creators WHERE order_id=?", (oid,))
    for eid in creators:
        db.execute("INSERT INTO order_creators (order_id, employee_id) VALUES (?,?)", (oid, eid))


@app.post("/api/orders")
def create_order():
    user = require_login()
    db = get_db()
    data, errors = read_form(db, user)
    if errors:
        return jsonify(error="; ".join(errors)), 400
    year = int(data["date"][:4])
    saved = []
    db.execute("BEGIN IMMEDIATE")
    try:
        if data["number"]:
            number = with_suffix(data["type"], data["number"])
            seq = seq_of(number)
        else:
            seq, number = next_number(db, data["type"], year)
        if duplicate_exists(db, data["type"], year, number):
            db.execute("ROLLBACK")
            return jsonify(error=f"Наказ № {number} за {year} рік у цьому реєстрі вже існує"), 409
        ts = now_iso()
        cur = db.execute(
            "INSERT INTO orders (type, year, seq, number, date, title, responsible, status,"
            " signer_id, created_by, updated_by, created_at, updated_at)"
            " VALUES (?,?,?,?,?,?,'',?,?,?,?,?,?)",
            (data["type"], year, seq, number, data["date"], data["title"], data["status"],
             data["signer_id"], user["id"], user["id"], ts, ts))
        oid = cur.lastrowid
        set_creators(db, oid, data["creators"])
        set_links(db, oid, data["links"])
        add_files(db, oid, request.files.getlist("files"), data["date"], saved)
        snap = snapshot(db, oid)
        audit(db, "create", snap, describe_new(snap))
        db.execute("COMMIT")
    except Exception:
        db.execute("ROLLBACK")
        for st in saved:            # файли, які не потрапили в базу
            remove_stored(st)
        raise
    return jsonify(fetch_order(db, oid, user)), 201


@app.put("/api/orders/<int:oid>")
def update_order(oid):
    user = require_login()
    db = get_db()
    old = db.execute("SELECT * FROM orders WHERE id=? AND deleted_at IS NULL", (oid,)).fetchone()
    if not old:
        abort(404)
    if not can_edit(user, old["type"]):
        abort(403)
    data, errors = read_form(db, user, fixed_type=old["type"])
    if errors:
        return jsonify(error="; ".join(errors)), 400
    year = int(data["date"][:4])
    number = data["number"] or old["number"]
    if number != old["number"]:     # старі номери не чіпаємо, суфікс лише для нововписаного
        number = with_suffix(old["type"], number)
    if duplicate_exists(db, old["type"], year, number, oid):
        return jsonify(error=f"Наказ № {number} за {year} рік у цьому реєстрі вже існує"), 409
    db.execute("BEGIN IMMEDIATE")
    to_delete, saved = [], []
    try:
        before = snapshot(db, oid)
        db.execute(
            "UPDATE orders SET year=?, seq=?, number=?, date=?, title=?, status=?,"
            " signer_id=?, updated_by=?, updated_at=? WHERE id=?",
            (year, seq_of(number), number, data["date"], data["title"], data["status"],
             data["signer_id"], user["id"], now_iso(), oid))
        set_creators(db, oid, data["creators"])
        set_links(db, oid, data["links"])
        rm = {int(x) for x in request.form.getlist("remove_files") if x.isdigit()}
        if rm:
            remove_files(db, oid, rm, defer=to_delete)
        add_files(db, oid, request.files.getlist("files"), data["date"], saved)
        after = snapshot(db, oid)
        changes = describe_changes(before, after)
        if changes:
            audit(db, "update", after, changes)
        db.execute("COMMIT")
    except Exception:
        db.execute("ROLLBACK")
        for st in saved:            # нові файли, які не потрапили в базу
            remove_stored(st)
        raise
    # диск чіпаємо лише після того, як база успішно збережена
    for st in to_delete:
        remove_stored(st)
    move_files(db, oid, data["date"])
    return jsonify(fetch_order(db, oid, user))


def check_action_password(user):
    """Пароль того, хто видаляє. Повертає відповідь-помилку або None, якщо пароль вірний."""
    key = ("del", user["id"])
    cnt, until = FAILS.get(key, (0, 0))
    if until > time.time():
        return jsonify(error=f"Забагато невдалих спроб. Спробуйте через {int((until - time.time()) / 60) + 1} хв."), 429
    pw = str((request.get_json(silent=True) or {}).get("password") or "")
    if not pw or not check_password_hash(user["password_hash"], pw):
        cnt += 1
        FAILS[key] = (0, time.time() + 120) if cnt >= 8 else (cnt, 0)
        return jsonify(error="Невірний пароль. Нічого не видалено."), 403
    FAILS.pop(key, None)
    return None


@app.delete("/api/orders/<int:oid>")
def delete_order(oid):
    """Видалення = переміщення в кошик: наказ і файли лишаються, адміністратор може відновити."""
    user = require_login()
    db = get_db()
    row = db.execute("SELECT type FROM orders WHERE id=? AND deleted_at IS NULL", (oid,)).fetchone()
    if not row:
        abort(404)
    if not can_edit(user, row["type"]):
        abort(403)
    err = check_action_password(user)
    if err:
        return err
    snap = snapshot(db, oid)
    db.execute("BEGIN IMMEDIATE")
    try:
        db.execute("UPDATE orders SET deleted_at=?, deleted_by=? WHERE id=?", (now_iso(), user["id"], oid))
        audit(db, "delete", snap, describe_new(snap))
        db.execute("COMMIT")
    except Exception:
        db.execute("ROLLBACK")
        raise
    return jsonify(ok=True)


UNDO_MINUTES = 10   # скільки хвилин той, хто видалив, може сам повернути наказ («Повернути» у повідомленні)


@app.post("/api/orders/<int:oid>/undo-delete")
def undo_delete(oid):
    user = require_login()
    db = get_db()
    row = db.execute("SELECT * FROM orders WHERE id=? AND deleted_at IS NOT NULL", (oid,)).fetchone()
    if not row:
        abort(404)
    fresh = row["deleted_at"] >= (dt.datetime.now() - dt.timedelta(minutes=UNDO_MINUTES)).strftime("%Y-%m-%d %H:%M:%S")
    if not user["is_admin"] and not (row["deleted_by"] == user["id"] and fresh and can_edit(user, row["type"])):
        return jsonify(error="Повернути вже не можна — зверніться до адміністратора (Кошик)"), 403
    return restore_order(db, row)


@app.post("/api/orders/<int:oid>/files")
def attach_files(oid):
    """Додати файли до наказу без відкриття форми (файл перетягнуто на рядок таблиці)."""
    user = require_login()
    db = get_db()
    row = db.execute("SELECT * FROM orders WHERE id=? AND deleted_at IS NULL", (oid,)).fetchone()
    if not row:
        abort(404)
    if not can_edit(user, row["type"]):
        abort(403)
    files = request.files.getlist("files")
    if not files:
        return jsonify(error="Немає файлів"), 400
    saved = []
    db.execute("BEGIN IMMEDIATE")
    try:
        add_files(db, oid, files, row["date"], saved)
        db.execute("UPDATE orders SET updated_by=?, updated_at=? WHERE id=?", (user["id"], now_iso(), oid))
        audit(db, "update", row, "Файли: " + ", ".join("+ " + (f.filename or "файл") for f in files))
        db.execute("COMMIT")
    except Exception:
        db.execute("ROLLBACK")
        for st in saved:
            remove_stored(st)
        raise
    return jsonify(fetch_order(db, oid, user))


# ---------------------------------------------------------------- Кошик (адміністратор)
def purge_order(db, oid):
    """Остаточно прибирає наказ із бази. Повертає файли, які треба стерти з диска після COMMIT."""
    stored = [r["stored"] for r in db.execute("SELECT stored FROM order_files WHERE order_id=?", (oid,))]
    db.execute("DELETE FROM orders WHERE id=?", (oid,))   # файли, автори, зв'язки — каскадом
    return stored


_last_purge = [0.0]


def purge_expired():
    """Накази, що пролежали в кошику понад TRASH_DAYS днів, видаляються остаточно."""
    cutoff = (dt.datetime.now() - dt.timedelta(days=TRASH_DAYS)).strftime("%Y-%m-%d %H:%M:%S")
    conn = connect()
    try:
        rows = conn.execute("SELECT id FROM orders WHERE deleted_at IS NOT NULL AND deleted_at < ?",
                            (cutoff,)).fetchall()
        for r in rows:
            snap = snapshot(conn, r["id"])
            conn.execute("BEGIN IMMEDIATE")
            try:
                stored = purge_order(conn, r["id"])
                audit(conn, "purge_auto", snap, f"Минуло {TRASH_DAYS} днів у кошику",
                      user={"id": None, "name": "Система"})
                conn.execute("COMMIT")
            except Exception:
                conn.execute("ROLLBACK")
                raise
            for st in stored:
                remove_stored(st)
    finally:
        conn.close()


@app.before_request
def housekeeping():
    if time.time() - _last_purge[0] > 3600:   # не частіше разу на годину
        _last_purge[0] = time.time()
        try:
            purge_expired()
        except Exception:
            app.logger.exception("Не вдалося очистити кошик")


@app.get("/api/trash")
def trash_list():
    require_admin()
    db = get_db()
    rows = db.execute(
        "SELECT o.id, o.type, o.number, o.date, o.title, o.deleted_at, e.name AS deleted_by_name, "
        "(SELECT COUNT(*) FROM order_files f WHERE f.order_id=o.id) AS files FROM orders o "
        "LEFT JOIN employees e ON e.id=o.deleted_by WHERE o.deleted_at IS NOT NULL "
        "ORDER BY o.deleted_at DESC").fetchall()
    out = []
    for r in rows:
        until = dt.datetime.strptime(r["deleted_at"], "%Y-%m-%d %H:%M:%S") + dt.timedelta(days=TRASH_DAYS)
        out.append({**dict(r), "reg": regs().get(r["type"], {}).get("short", r["type"]),
                    "purge_on": until.strftime("%Y-%m-%d")})
    return jsonify(orders=out, days=TRASH_DAYS)


@app.post("/api/trash/<int:oid>/restore")
def trash_restore(oid):
    require_admin()
    db = get_db()
    row = db.execute("SELECT * FROM orders WHERE id=? AND deleted_at IS NOT NULL", (oid,)).fetchone()
    if not row:
        abort(404)
    return restore_order(db, row)


def restore_order(db, row):
    oid = row["id"]
    if row["type"] not in regs():
        return jsonify(error="Реєстру цього наказу вже немає"), 409
    dup = duplicate_exists(db, row["type"], row["year"], row["number"], oid)
    if dup:
        return jsonify(error=f"Номер {row['number']} за {row['year']} рік уже зайняв інший наказ "
                             f"(від {fmt_d(dup['date'])}). Змініть номер того наказу, потім відновіть цей."), 409
    db.execute("BEGIN IMMEDIATE")
    try:
        db.execute("UPDATE orders SET deleted_at=NULL, deleted_by=NULL WHERE id=?", (oid,))
        audit(db, "restore", row, "Відновлено з кошика")
        db.execute("COMMIT")
    except Exception:
        db.execute("ROLLBACK")
        raise
    return jsonify(ok=True)


@app.delete("/api/trash/<int:oid>")
def trash_purge(oid):
    user = require_admin()
    db = get_db()
    if not db.execute("SELECT 1 FROM orders WHERE id=? AND deleted_at IS NOT NULL", (oid,)).fetchone():
        abort(404)
    err = check_action_password(user)
    if err:
        return err
    snap = snapshot(db, oid)
    db.execute("BEGIN IMMEDIATE")
    try:
        stored = purge_order(db, oid)
        audit(db, "purge", snap, describe_new(snap))
        db.execute("COMMIT")
    except Exception:
        db.execute("ROLLBACK")
        raise
    for st in stored:
        remove_stored(st)
    return jsonify(ok=True)


def _file_row(fid):
    user = current_user()
    row = get_db().execute("SELECT f.name, f.stored, o.type, o.deleted_at FROM order_files f "
                           "JOIN orders o ON o.id=f.order_id WHERE f.id=?", (fid,)).fetchone()
    if not row or not can_view(user, row["type"]) or (row["deleted_at"] and not (user and user["is_admin"])):
        abort(404)
    if not can_open_files(user):
        abort(401)
    return row


@app.get("/api/files/<int:fid>")
def download_file(fid):
    row = _file_row(fid)
    path = stored_path(row["stored"])
    if not path or not os.path.exists(path):
        abort(404)
    ext = file_ext(row["name"], row["stored"])
    # у браузері відкриваємо лише безпечні типи; .html, .svg та все інше — лише завантаженням
    inline = ext in DIRECT_EXT and request.args.get("download") != "1"
    resp = send_file(path, download_name=row["name"], as_attachment=not inline,
                     mimetype=SAFE_MIME.get(ext, "application/octet-stream"))
    resp.headers["X-Content-Type-Options"] = "nosniff"
    if ext not in (".pdf",):
        resp.headers["Content-Security-Policy"] = "default-src 'none'; style-src 'unsafe-inline'; sandbox"
    return resp


@app.get("/api/files/<int:fid>/preview")
def preview_file(fid):
    """Файл для вбудованого перегляду: PDF/зображення як є, Word і Excel — у вигляді PDF."""
    row = _file_row(fid)
    ext = file_ext(row["name"], row["stored"])
    if ext in OFFICE_EXT:
        pdf = preview_pdf(row["stored"], ext)
        if not pdf:
            return Response("Не вдалося підготувати перегляд документа. Завантажте файл кнопкою «⬇».",
                            503, mimetype="text/plain")
        resp = send_file(pdf, mimetype="application/pdf",
                         download_name=os.path.splitext(row["name"] or "document")[0] + ".pdf",
                         as_attachment=False)
        resp.headers["X-Content-Type-Options"] = "nosniff"
        return resp
    if ext in IMG_EXT:
        return image_page(f"/api/files/{fid}", row["name"])
    if ext in DIRECT_EXT:
        return download_file(fid)
    if ext in ARCH_EXT:
        return archive_page(fid, row)
    abort(415)


IMG_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}


def image_page(src, name):
    """Зображення для перегляду: на всю ширину панелі; клік — оригінал у новій вкладці."""
    import html
    page = f"""<!doctype html><html lang="uk"><head><meta charset="utf-8"><title>{html.escape(name or "зображення")}</title>
<style>html,body{{margin:0;background:#e9ebef}} img{{display:block;width:100%;height:auto}} a{{display:block;cursor:zoom-in}}</style>
</head><body><a href="{html.escape(src)}" target="_blank" rel="noopener" title="Відкрити в повному розмірі"><img src="{html.escape(src)}" alt=""></a></body></html>"""
    resp = Response(page, mimetype="text/html")
    resp.headers["Content-Security-Policy"] = "default-src 'none'; img-src 'self'; style-src 'unsafe-inline'; frame-ancestors 'self'"
    return resp


# ---------------------------------------------------------------- Архіви: перелік файлів і перегляд окремого файлу
ARCH_EXT = {".zip", ".7z", ".rar", ".tar", ".tgz", ".gz"}
ARCH_MAX_FILE = 300 * 1024 ** 2   # більший файл з архіву не розпаковуємо для перегляду
ARCH_MAX_LIST = 5000


class ArchiveError(Exception):
    pass


# розпаковка одного файлу з 7Z в окремому процесі: argv = архів, куди, ім'я файлу, ліміт розміру в байтах
SEVENZ_EXTRACT = """
import sys
try:
    import resource
    lim = int(sys.argv[4])
    resource.setrlimit(resource.RLIMIT_FSIZE, (lim, lim))
except ImportError:
    pass   # не Linux (лише для розробки) — без ліміту
import py7zr
with py7zr.SevenZipFile(sys.argv[1]) as z:
    z.extract(path=sys.argv[2], targets=[sys.argv[3]])
"""


def _zip_name(info):
    """Імена з ZIP, створених у Windows без позначки UTF-8, записані в cp866 — інакше кирилиця «ламається»."""
    name = info.filename
    if not info.flag_bits & 0x800:
        try:
            raw = name.encode("cp437")
        except UnicodeEncodeError:
            return name
        for enc in ("utf-8", "cp866"):   # Linux/macOS пишуть UTF-8 без позначки, Windows — cp866
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                continue
    return name


def _libarchive():
    try:
        import libarchive   # потребує системної бібліотеки libarchive (є в Docker-образі)
        return libarchive
    except (ImportError, OSError):
        raise ArchiveError("Формат RAR на цьому сервері не підтримується")


def arch_entries(path, ext):
    """[(назва, розмір, це_папка)] у сталому порядку; номер у списку — адреса файлу в архіві."""
    try:
        if ext == ".zip":
            import zipfile
            with zipfile.ZipFile(path) as z:
                return [(_zip_name(i), i.file_size, i.is_dir()) for i in z.infolist()][:ARCH_MAX_LIST]
        if ext == ".7z":
            import py7zr
            with py7zr.SevenZipFile(path) as z:
                return [(i.filename, i.uncompressed or 0, i.is_directory) for i in z.list()][:ARCH_MAX_LIST]
        if ext == ".rar":
            libarchive = _libarchive()
            out = []
            with libarchive.file_reader(path) as a:
                for e in a:
                    out.append((e.pathname, e.size or 0, e.isdir))
                    if len(out) >= ARCH_MAX_LIST:
                        break
            return out
        import tarfile
        with tarfile.open(path) as t:
            return [(m.name, m.size, m.isdir()) for m in t.getmembers() if m.isfile() or m.isdir()][:ARCH_MAX_LIST]
    except ArchiveError:
        raise
    except ImportError:
        raise ArchiveError("Цей формат архіву на сервері не підтримується")
    except Exception as e:   # пошкоджений, зашифрований чи незнайомий архів
        msg = str(e).lower()
        raise ArchiveError("Архів захищено паролем" if "password" in msg or "encrypt" in msg
                           else "Не вдалося прочитати архів (можливо, пошкоджений)")


def arch_read(path, ext, n):
    """Вміст n-го файлу архіву (bytes)."""
    try:
        if ext == ".zip":
            import zipfile
            with zipfile.ZipFile(path) as z:
                return z.open(z.infolist()[n]).read(ARCH_MAX_FILE + 1)
        if ext == ".7z":
            import py7zr
            with py7zr.SevenZipFile(path) as z:
                name = z.list()[n].filename
            tmp = tempfile.mkdtemp(prefix="arch_", dir=PREVIEWS)
            try:
                # py7zr не вміє зупинити розпаковку на заданому розмірі, тож розпаковуємо в окремому процесі
                # з системним обмеженням розміру файлу: більший за ліміт файл (хоч би що казав заголовок архіву)
                # не допише себе на диск — процес буде зупинено
                r = subprocess.run([sys.executable, "-c", SEVENZ_EXTRACT, path, tmp, name, str(ARCH_MAX_FILE + 1)],
                                   capture_output=True, timeout=180)
                if r.returncode != 0:
                    raise ArchiveError("Файл в архіві завеликий для перегляду або архів пошкоджений — завантажте архів цілком")
                f = os.path.normpath(os.path.join(tmp, name))
                if os.path.commonpath([f, tmp]) != os.path.normpath(tmp):
                    raise ArchiveError("Недопустима назва файлу в архіві")
                with open(f, "rb") as fh:
                    return fh.read(ARCH_MAX_FILE + 1)
            finally:
                shutil.rmtree(tmp, ignore_errors=True)
        if ext == ".rar":
            libarchive = _libarchive()
            with libarchive.file_reader(path) as a:
                for i, e in enumerate(a):
                    if i == n:
                        data = bytearray()
                        for block in e.get_blocks():
                            data += block
                            if len(data) > ARCH_MAX_FILE:
                                break
                        return bytes(data)
            raise ArchiveError("Файл в архіві не знайдено")
        import tarfile
        with tarfile.open(path) as t:
            members = [m for m in t.getmembers() if m.isfile() or m.isdir()]
            return t.extractfile(members[n]).read(ARCH_MAX_FILE + 1)
    except ArchiveError:
        raise
    except Exception:
        raise ArchiveError("Не вдалося розпакувати файл з архіву")


def human_size(n):
    for unit in ("Б", "КБ", "МБ", "ГБ"):
        if n < 1024 or unit == "ГБ":
            return f"{n:.0f} {unit}" if unit == "Б" else f"{n:.1f} {unit}".replace(".", ",")
        n /= 1024


def archive_page(fid, row):
    """HTML-сторінка для перегляду: перелік файлів архіву з посиланнями «переглянути» / «завантажити»."""
    import html
    path = stored_path(row["stored"])
    ext = file_ext(row["name"], row["stored"])
    try:
        entries = arch_entries(path, ext)
        err = ""
    except ArchiveError as e:
        entries, err = [], str(e)
    viewable = DIRECT_EXT | OFFICE_EXT
    rows = []
    for i, (name, size, is_dir) in enumerate(entries):
        if is_dir:
            continue
        iext = os.path.splitext(name)[1].lower()
        acts = []
        if iext in viewable and size <= ARCH_MAX_FILE:
            acts.append(f'<a href="/api/files/{fid}/entry?n={i}&view=1" target="_blank" rel="noopener">переглянути</a>')
        acts.append(f'<a href="/api/files/{fid}/entry?n={i}">завантажити</a>')
        folder, base = os.path.split(name.rstrip("/"))
        rows.append(f'<tr><td><div>{"<span class=d>" + html.escape(folder) + "/</span>" if folder else ""}{html.escape(base)}</div>'
                    f'<div class=m>{human_size(size)} · {" · ".join(acts)}</div></td></tr>')
    title = html.escape(row["name"] or "архів")
    body = (f'<p class=e>{html.escape(err)}. Архів можна завантажити кнопкою «⬇».</p>' if err else
            f'<table>{"".join(rows)}</table>' if rows else '<p class=e>Архів порожній.</p>')
    more = f'<p class=d>Показано перші {ARCH_MAX_LIST} записів.</p>' if len(entries) >= ARCH_MAX_LIST else ""
    page = f"""<!doctype html><html lang="uk"><head><meta charset="utf-8"><title>{title}</title><style>
body{{font:14px/1.45 system-ui,"Segoe UI",Roboto,sans-serif;margin:0;padding:14px 16px;color:#1d2330;background:#fff}}
h1{{font-size:15px;margin:0 0 10px}} h1 span{{color:#6b7385;font-weight:400}}
table{{width:100%;border-collapse:collapse}} td{{padding:6px 8px;border-bottom:1px solid #e3e6ec;vertical-align:top;overflow-wrap:anywhere}}
tr:nth-child(even) td{{background:#f7f8fa}} .d{{color:#6b7385}} .m{{font-size:12.5px;color:#6b7385;margin-top:2px}}
a{{color:#2f5bea;text-decoration:none}} a:hover{{text-decoration:underline}} .e{{color:#c93a3a}}
</style></head><body><h1>🗜 {title} <span>· файлів: {len(rows)}</span></h1>{body}{more}</body></html>"""
    resp = Response(page, mimetype="text/html")
    resp.headers["Content-Security-Policy"] = "default-src 'none'; style-src 'unsafe-inline'; frame-ancestors 'self'"
    return resp


@app.get("/api/files/<int:fid>/entry")
def archive_entry(fid):
    """Окремий файл з архіву: переглянути (як звичайне вкладення) або завантажити."""
    row = _file_row(fid)
    ext = file_ext(row["name"], row["stored"])
    n = request.args.get("n", "")
    if ext not in ARCH_EXT or not n.isdigit():
        abort(404)
    n = int(n)
    path = stored_path(row["stored"])
    try:
        entries = arch_entries(path, ext)
        if n >= len(entries) or entries[n][2]:
            abort(404)
        name, size, _ = entries[n]
        if size > ARCH_MAX_FILE:
            return Response("Файл завеликий для перегляду — завантажте архів цілком.", 413, mimetype="text/plain")
        iext = os.path.splitext(name)[1].lower()
        # розпакований файл кешується поряд із PDF-переглядами (видаляється разом з архівом)
        cached = os.path.join(PREVIEWS, f"{os.path.basename(row['stored'])}__{n}{re.sub(r'[^a-z0-9.]', '', iext)[:10]}")
        if not (os.path.exists(cached) and os.path.getmtime(cached) >= os.path.getmtime(path)):
            data = arch_read(path, ext, n)
            if len(data) > ARCH_MAX_FILE:
                return Response("Файл завеликий для перегляду — завантажте архів цілком.", 413, mimetype="text/plain")
            with open(cached + ".tmp", "wb") as f:
                f.write(data)
            os.replace(cached + ".tmp", cached)
    except ArchiveError as e:
        return Response(f"{e}. Завантажте архів кнопкою «⬇».", 422, mimetype="text/plain")
    base = os.path.basename(name.rstrip("/")) or "file"
    view = request.args.get("view") == "1"
    if view and iext in IMG_EXT and request.args.get("raw") != "1":
        return image_page(f"/api/files/{fid}/entry?n={n}&view=1&raw=1", base)
    if view and iext in OFFICE_EXT:
        pdf = convert_pdf(cached, iext, cached + ".pdf")
        if not pdf:
            return Response("Не вдалося підготувати перегляд документа. Завантажте файл.", 503,
                            mimetype="text/plain")
        resp = send_file(pdf, mimetype="application/pdf", download_name=os.path.splitext(base)[0] + ".pdf")
    else:
        inline = view and iext in DIRECT_EXT
        resp = send_file(cached, download_name=base, as_attachment=not inline,
                         mimetype=SAFE_MIME.get(iext, "application/octet-stream"))
        if iext != ".pdf":
            resp.headers["Content-Security-Policy"] = "default-src 'none'; style-src 'unsafe-inline'; sandbox"
    resp.headers["X-Content-Type-Options"] = "nosniff"
    return resp


# ---------------------------------------------------------------- Довідники (адміністратор)
@app.get("/api/lookups")
def lookups():
    require_login()
    db = get_db()
    return jsonify(
        # службовий запис «Адміністратор» у списки вибору («Хто створив», «Хто підписав») не потрапляє
        employees=[dict(r) for r in db.execute(
            "SELECT id, name, active FROM employees WHERE name<>'Адміністратор' ORDER BY name")],
        signers=[dict(r) for r in db.execute(
            "SELECT id, name, position, active FROM signers WHERE name<>'Адміністратор' ORDER BY name")])


def flag(d, key, default=0):
    v = d.get(key, default)
    return 1 if v in (True, 1, "1", "true", "on") else 0


def admin_remains(db):
    return db.execute("SELECT COUNT(*) FROM employees WHERE is_admin=1 AND active=1 "
                      "AND password_hash IS NOT NULL").fetchone()[0] > 0


def emp_json(r):
    return {"id": r["id"], "name": r["name"], "active": bool(r["active"]),
            "is_admin": bool(r["is_admin"]), "is_secretary": bool(r["is_secretary"]),
            "is_hr": bool(r["is_hr"]), "is_viewer": bool(r["is_viewer"]), "has_password": bool(r["password_hash"]),
            "used": r["used"]}


ROLE_NAMES = ("Адміністратор", "Секретар", "Кадровик", "Переглядач")


def role_text(roles):
    return ", ".join(n for n, r in zip(ROLE_NAMES, roles) if r) or "без ролей"


EMP_SELECT = ("SELECT e.*, (SELECT COUNT(*) FROM order_creators oc WHERE oc.employee_id=e.id) "
              "AS used FROM employees e")


@app.get("/api/employees")
def employees_list():
    require_admin()
    rows = get_db().execute(EMP_SELECT + " ORDER BY e.name").fetchall()
    return jsonify(employees=[emp_json(r) for r in rows])


def name_taken(db, table, name, exclude_id=0):
    key = norm(name).casefold()
    return any(norm(r["name"]).casefold() == key and r["id"] != exclude_id
               for r in db.execute(f"SELECT id, name FROM {table}"))


@app.post("/api/employees")
def employees_create():
    require_admin()
    db = get_db()
    d = request.get_json(silent=True) or {}
    name = norm(d.get("name"))
    if not name:
        return jsonify(error="Вкажіть ПІП"), 400
    if name_taken(db, "employees", name):
        return jsonify(error="Співробітник з таким ПІП уже є"), 409
    roles = (flag(d, "is_admin"), flag(d, "is_secretary"), flag(d, "is_hr"), flag(d, "is_viewer"))
    pw = str(d.get("password") or "")
    if any(roles) and len(pw) < MIN_PASSWORD:
        return jsonify(error=f"Для ролі потрібен пароль (мінімум {MIN_PASSWORD} символів)"), 400
    if pw and len(pw) < MIN_PASSWORD:
        return jsonify(error=f"Пароль має містити щонайменше {MIN_PASSWORD} символів"), 400
    cur = db.execute(
        "INSERT INTO employees (name, active, is_admin, is_secretary, is_hr, is_viewer, password_hash, "
        "created_at) VALUES (?,?,?,?,?,?,?,?)",
        (name, flag(d, "active", 1), *roles, generate_password_hash(pw) if pw else None,
         now_iso()))
    audit(db, "emp_create", details=f"{name}; ролі: {role_text(roles)}" + ("; пароль задано" if pw else ""))
    return jsonify(emp_json(db.execute(EMP_SELECT + " WHERE e.id=?", (cur.lastrowid,)).fetchone())), 201


@app.put("/api/employees/<int:eid>")
def employees_update(eid):
    me = require_admin()
    db = get_db()
    old = db.execute("SELECT * FROM employees WHERE id=?", (eid,)).fetchone()
    if not old:
        abort(404)
    d = request.get_json(silent=True) or {}
    name = norm(d.get("name", old["name"]))
    if not name:
        return jsonify(error="Вкажіть ПІП"), 400
    if name_taken(db, "employees", name, eid):
        return jsonify(error="Співробітник з таким ПІП уже є"), 409
    roles = (flag(d, "is_admin", old["is_admin"]), flag(d, "is_secretary", old["is_secretary"]),
             flag(d, "is_hr", old["is_hr"]), flag(d, "is_viewer", old["is_viewer"]))
    active = flag(d, "active", old["active"])
    pw = str(d.get("password") or "")
    if pw and len(pw) < MIN_PASSWORD:
        return jsonify(error=f"Пароль має містити щонайменше {MIN_PASSWORD} символів"), 400
    pw_hash = generate_password_hash(pw) if pw else old["password_hash"]
    if any(roles) and not pw_hash:
        return jsonify(error="Для ролі потрібен пароль — задайте його"), 400
    if eid == me["id"] and (not roles[0] or not active):
        return jsonify(error="Не можна забрати в себе роль адміністратора або вимкнути себе"), 409
    db.execute("BEGIN IMMEDIATE")
    try:
        db.execute("UPDATE employees SET name=?, active=?, is_admin=?, is_secretary=?, is_hr=?, is_viewer=?, "
                   "password_hash=?, must_change=? WHERE id=?",
                   (name, active, *roles, pw_hash, 0 if pw else old["must_change"], eid))
        if not admin_remains(db):
            db.execute("ROLLBACK")
            return jsonify(error="Має залишитись хоча б один активний адміністратор із паролем"), 409
        old_roles = (old["is_admin"], old["is_secretary"], old["is_hr"], old["is_viewer"])
        ch = []
        if name != old["name"]:
            ch.append(f"ПІП: «{old['name']}» → «{name}»")
        if tuple(map(bool, roles)) != tuple(map(bool, old_roles)):
            ch.append(f"Ролі: {role_text(old_roles)} → {role_text(roles)}")
        if bool(active) != bool(old["active"]):
            ch.append("Активний: " + ("так" if active else "ні"))
        if pw:
            ch.append("Пароль змінено")
        if ch:
            audit(db, "emp_update", details=[name] + ch)
        db.execute("COMMIT")
    except Exception:
        db.execute("ROLLBACK")
        raise
    if eid == me["id"] and pw:
        session["ph"] = pw_stamp(pw_hash)   # власний пароль змінено в довіднику — поточний сеанс лишається
    return jsonify(emp_json(db.execute(EMP_SELECT + " WHERE e.id=?", (eid,)).fetchone()))


@app.delete("/api/employees/<int:eid>")
def employees_delete(eid):
    me = require_admin()
    db = get_db()
    row = db.execute(EMP_SELECT + " WHERE e.id=?", (eid,)).fetchone()
    if not row:
        abort(404)
    if eid == me["id"]:
        return jsonify(error="Не можна видалити самого себе"), 409
    if row["used"]:
        return jsonify(error=f"Прив'язаний до наказів ({row['used']}) — зніміть «Активний» "
                             "замість видалення"), 409
    db.execute("DELETE FROM employees WHERE id=?", (eid,))
    audit(db, "emp_delete", details=row["name"])
    return jsonify(ok=True)


@app.get("/api/signers")
def signers_list():
    require_admin()
    rows = get_db().execute(
        "SELECT s.*, (SELECT COUNT(*) FROM orders o WHERE o.signer_id=s.id) AS used "
        "FROM signers s ORDER BY s.name").fetchall()
    return jsonify(signers=[{"id": r["id"], "name": r["name"], "position": r["position"],
                             "active": bool(r["active"]), "used": r["used"]} for r in rows])


@app.post("/api/signers")
def signers_create():
    require_admin()
    db = get_db()
    d = request.get_json(silent=True) or {}
    name = norm(d.get("name"))
    if not name:
        return jsonify(error="Вкажіть ПІП"), 400
    if name_taken(db, "signers", name):
        return jsonify(error="Такий підписант уже є"), 409
    cur = db.execute("INSERT INTO signers (name, position, active) VALUES (?,?,?)",
                     (name, norm(d.get("position")), flag(d, "active", 1)))
    audit(db, "sig_create", details=name)
    return jsonify(id=cur.lastrowid), 201


@app.put("/api/signers/<int:sid>")
def signers_update(sid):
    require_admin()
    db = get_db()
    old = db.execute("SELECT * FROM signers WHERE id=?", (sid,)).fetchone()
    if not old:
        abort(404)
    d = request.get_json(silent=True) or {}
    name = norm(d.get("name", old["name"]))
    if not name:
        return jsonify(error="Вкажіть ПІП"), 400
    if name_taken(db, "signers", name, sid):
        return jsonify(error="Такий підписант уже є"), 409
    position, active = norm(d.get("position", old["position"])), flag(d, "active", old["active"])
    db.execute("UPDATE signers SET name=?, position=?, active=? WHERE id=?", (name, position, active, sid))
    ch = ([f"ПІП: «{old['name']}» → «{name}»"] if name != old["name"] else []) + \
         ([f"Посада: «{old['position']}» → «{position}»"] if position != old["position"] else []) + \
         (["Активний: " + ("так" if active else "ні")] if bool(active) != bool(old["active"]) else [])
    if ch:
        audit(db, "sig_update", details=[name] + ch)
    return jsonify(ok=True)


@app.delete("/api/signers/<int:sid>")
def signers_delete(sid):
    require_admin()
    db = get_db()
    used = db.execute("SELECT COUNT(*) FROM orders WHERE signer_id=?", (sid,)).fetchone()[0]
    if used:
        return jsonify(error=f"Підписант указаний у наказах ({used}) — зніміть «Активний» "
                             "замість видалення"), 409
    row = db.execute("SELECT name FROM signers WHERE id=?", (sid,)).fetchone()
    db.execute("DELETE FROM signers WHERE id=?", (sid,))
    if row:
        audit(db, "sig_delete", details=row["name"])
    return jsonify(ok=True)


# ---------------------------------------------------------------- Налаштування (адміністратор)
@app.get("/api/settings")
def settings_get():
    require_admin()
    return jsonify(public_access=public_access(), options=PUBLIC_ACCESS,
                   ca=ca_info(), ca_public=ca_setup_phase(), cloud=cloud_status(),
                   backup={**backup_schedule(), "pending": os.path.exists(os.path.join(DATA_DIR, "backup_request")),
                           "log": cloud_log()})


CLOUD_STALE_DAYS = 2   # копія в хмару давніша — попередження адміністратору в «Перевірці»
CLOUD_RUN_MAX_HOURS = 12   # довше «виконується» не буває навіть для першої повної копії


def cloud_status():
    """Стан копії в хмару (пише backup_gdrive.sh на сервері); None — копію не налаштовано."""
    import json
    try:
        with open(os.path.join(DATA_DIR, "cloud_backup.json"), encoding="utf-8") as f:
            st = json.load(f)
        t = dt.datetime.strptime(st["time"], "%Y-%m-%d %H:%M:%S")
    except (OSError, ValueError, KeyError, TypeError):
        return None
    age = (dt.datetime.now() - t).total_seconds() / 86400
    ok, msg, running = bool(st.get("ok")), str(st.get("message", "")), bool(st.get("running"))
    if running and age > CLOUD_RUN_MAX_HOURS / 24:   # «виконується» надто довго — копію перервано (перезавантаження тощо)
        ok, running, msg = False, False, "копію перервано (сервер перезавантажувався або зник зв'язок) — запустіть ще раз"
    return {"time": st["time"], "ok": ok, "message": msg, "stale": age > CLOUD_STALE_DAYS, "running": running}


def backup_schedule():
    t = setting("backup_time", "02:30")
    return {"time": t if re.fullmatch(r"\d\d:\d\d", t) else "02:30", "enabled": setting("backup_enabled", "1") == "1"}


def write_backup_schedule(conn=None):
    """Розклад для сервера: backup_gdrive.sh --auto читає data/backup_schedule.json (час — за TIMEZONE)."""
    import json
    if conn is not None:
        st = {r["key"]: r["value"] for r in conn.execute("SELECT key, value FROM settings")}
        t, en = st.get("backup_time", "02:30"), st.get("backup_enabled", "1") == "1"
    else:
        b = backup_schedule()
        t, en = b["time"], b["enabled"]
    try:
        with open(os.path.join(DATA_DIR, "backup_schedule.json"), "w", encoding="utf-8") as f:
            json.dump({"time": t, "enabled": en}, f)
    except OSError:
        pass


def cloud_log():
    """Останні рядки журналу останньої копії (пише backup_gdrive.sh)."""
    try:
        with open(os.path.join(DATA_DIR, "cloud_backup.log"), encoding="utf-8", errors="replace") as f:
            return "".join(f.readlines()[-40:])
    except OSError:
        return ""


# ---------------------------------------------------------------- Сертифікат HTTPS (кореневий сертифікат Caddy)
# Caddy віддає свій кореневий сертифікат програмі через внутрішню мережу Docker (див. Caddyfile, :8081).
CA_URL = os.environ.get("CA_URL", "")
_ca_cache = {"t": 0.0, "pem": None}


def ca_pem():
    if not (BEHIND_PROXY and CA_URL):
        return None
    if time.time() - _ca_cache["t"] > 300 or not _ca_cache["pem"]:   # кеш на 5 хв
        import urllib.request
        try:
            with urllib.request.urlopen(CA_URL, timeout=3) as r:
                pem = r.read(65536).decode("ascii", "replace")
            _ca_cache["pem"] = pem if "BEGIN CERTIFICATE" in pem else None
        except OSError:
            _ca_cache["pem"] = None
        _ca_cache["t"] = time.time()
    return _ca_cache["pem"]


def ca_info():
    """Відбитки кореневого сертифіката для звірки на комп'ютері; None — HTTPS не ввімкнено."""
    pem = ca_pem()
    if not pem:
        return None
    import hashlib
    import ssl
    try:
        der = ssl.PEM_cert_to_DER_cert(pem)
    except ValueError:
        return None
    fp = lambda h: ":".join(f"{b:02X}" for b in h(der).digest())
    return {"sha1": fp(hashlib.sha1), "sha256": fp(hashlib.sha256)}


def ca_setup_phase():
    """Перше налаштування: адміністратор ще не змінив початковий пароль — сертифікат можна взяти без входу."""
    return bool(get_db().execute("SELECT 1 FROM employees WHERE is_admin=1 AND active=1 AND must_change=1 "
                                 "AND password_hash IS NOT NULL").fetchone())


def ca_file_response():
    pem = ca_pem()
    if not pem:
        abort(404)
    return Response(pem, mimetype="application/x-x509-ca-cert",
                    headers={"Content-Disposition": "attachment; filename=nakazy-ca.crt"})


@app.get("/nakazy-ca.crt")
def ca_public():
    """Посилання для першого комп'ютера (працює й по HTTP); після першого налаштування — лише для адміністратора."""
    u = current_user()
    if not ca_pem():
        abort(404)
    if not ca_setup_phase() and not (u and u["is_admin"]):
        return Response("Сертифікат видає адміністратор: «Адміністрування → Налаштування → Сертифікат HTTPS».",
                        403, mimetype="text/plain")
    return ca_file_response()


@app.get("/api/ca.crt")
def ca_admin():
    require_admin()
    return ca_file_response()


@app.put("/api/settings")
def settings_put():
    require_admin()
    d = request.get_json(silent=True) or {}
    db = get_db()
    changes = []
    if "public_access" in d:
        v = d.get("public_access")
        if v not in PUBLIC_ACCESS:
            return jsonify(error="Невідоме значення доступу"), 400
        old = public_access()
        if v != old:
            set_setting(db, "public_access", v)
            changes.append(f"Доступ без входу: «{PUBLIC_ACCESS[old]}» → «{PUBLIC_ACCESS[v]}»")
    if "backup_time" in d or "backup_enabled" in d:
        bt = str(d.get("backup_time", backup_schedule()["time"])).strip()
        if not re.fullmatch(r"([01]\d|2[0-3]):[0-5]\d", bt):
            return jsonify(error="Час копії — у форматі ГГ:ХХ, наприклад 02:30"), 400
        be = flag(d, "backup_enabled", 1 if backup_schedule()["enabled"] else 0)
        old = backup_schedule()
        if bt != old["time"] or bool(be) != old["enabled"]:
            set_setting(db, "backup_time", bt)
            set_setting(db, "backup_enabled", "1" if be else "0")
            changes.append(f"Копія в хмару: {'щодня о ' + bt if be else 'автоматичну копію вимкнено'}")
            write_backup_schedule()
    if changes:
        audit(db, "settings", details=changes)
    return jsonify(ok=True)


def set_setting(db, key, value):
    db.execute("INSERT INTO settings (key, value) VALUES (?, ?) "
               "ON CONFLICT(key) DO UPDATE SET value=excluded.value", (key, value))
    g.pop("settings", None)


@app.post("/api/backup/run")
def backup_run():
    """«Зробити копію зараз»: сервер (backup_gdrive.sh --auto) підхопить запит протягом кількох хвилин."""
    require_admin()
    with open(os.path.join(DATA_DIR, "backup_request"), "w") as f:
        f.write(now_iso())
    audit(get_db(), "settings", details="Копія в хмару: запущено вручну")
    return jsonify(ok=True)


# ---------------------------------------------------------------- Пошук наказу для зв'язку
@app.get("/api/order-search")
def order_search():
    """Накази за номером — для поля «Скасовує / Вносить зміни до» (лише реєстри того самого виду)."""
    user = require_login()
    num = norm(request.args.get("num", "")).lstrip("№# ").strip()
    if not num:
        return jsonify(orders=[])
    kind = regs().get(request.args.get("type", ""), {}).get("kind", "main")
    keys = [k for k, r in regs().items() if r["kind"] == kind and can_view(user, k)]
    if not keys:
        return jsonify(orders=[])
    sql = (f"SELECT id, type, number, date, title FROM orders WHERE type IN ({','.join('?' * len(keys))}) "
           "AND deleted_at IS NULL AND ")
    params = list(keys)
    if num.isdigit():
        sql += "(ULOWER(number)=? OR ULOWER(number) LIKE ? ESCAPE '\\')"
        params += [num, num + "-%"]
    else:
        sql += "ULOWER(number) LIKE ? ESCAPE '\\'"
        params.append(f"%{like_esc(num.casefold())}%")
    year = request.args.get("year", "").strip()
    if year.isdigit():
        sql += " AND year=?"
        params.append(int(year))
    rows = get_db().execute(sql + " ORDER BY date DESC, id DESC LIMIT 30", params).fetchall()
    return jsonify(orders=[{**dict(r), "reg": regs()[r["type"]]["short"]} for r in rows])


# ---------------------------------------------------------------- Журнал дій (адміністратор)
AUDIT_GROUPS = {
    "orders": ("create", "update", "delete", "restore", "purge", "purge_auto"),
    "trash": ("delete", "restore", "purge", "purge_auto"),
    "dirs": ("emp_create", "emp_update", "emp_delete", "sig_create", "sig_update", "sig_delete",
             "reg_create", "reg_update", "reg_delete"),
    "security": ("password", "lock", "emp_update", "settings"),
    "import": ("import",),
}


@app.get("/api/audit")
def audit_list():
    require_admin()
    where, params = [], []
    oid = request.args.get("order", "").strip()
    if oid.isdigit():
        where.append("a.order_id=?")
        params.append(int(oid))
    grp = AUDIT_GROUPS.get(request.args.get("group", ""))
    if grp:
        where.append(f"a.action IN ({','.join('?' * len(grp))})")
        params += grp
    q = norm(request.args.get("q", ""))
    if q:
        like = f"%{like_esc(q.casefold())}%"
        where.append("(ULOWER(a.details) LIKE ? ESCAPE '\\' OR ULOWER(a.user_name) LIKE ? ESCAPE '\\' "
                     "OR ULOWER(COALESCE(a.number,'')) LIKE ? ESCAPE '\\')")
        params += [like, like, like]
    before = request.args.get("before", "").strip()
    if before.isdigit():
        where.append("a.id<?")
        params.append(int(before))
    sql = "SELECT a.* FROM audit a" + (" WHERE " + " AND ".join(where) if where else "") + \
          " ORDER BY a.id DESC LIMIT 201"
    rows = [dict(r) for r in get_db().execute(sql, params)]
    for r in rows:
        r["reg_short"] = regs().get(r["reg"], {}).get("short", r["reg"] or "")
    return jsonify(entries=rows[:200], more=len(rows) > 200)


# ---------------------------------------------------------------- Перевірка реєстру
def compress_ranges(nums):
    """[3,4,5,9] -> ['3–5', '9']"""
    out, start, prev = [], None, None
    for n in nums:
        if start is None:
            start = prev = n
        elif n == prev + 1:
            prev = n
        else:
            out.append(str(start) if start == prev else f"{start}–{prev}")
            start = prev = n
    if start is not None:
        out.append(str(start) if start == prev else f"{start}–{prev}")
    return out


def check_periods(today=None):
    """Нагадування: раз на тиждень (перший вхід тижня) і в останні 5 днів місяця."""
    today = today or dt.date.today()
    iso = today.isocalendar()
    periods = [f"W{iso[0]}-{iso[1]:02d}"]
    nxt = (today.replace(day=28) + dt.timedelta(days=4)).replace(day=1)
    if (nxt - today).days <= 5:
        periods.append(f"M{today.year}-{today.month:02d}")
    return periods


@app.get("/api/checks")
def checks():
    user = require_login()
    if not (user["is_admin"] or user["is_secretary"] or user["is_hr"]):
        abort(403)
    db = get_db()
    today = dt.date.today()
    y = request.args.get("year", "").strip()
    years = [int(y)] if y.isdigit() else ([today.year - 1, today.year] if today.month == 1 else [today.year])
    stale = (today - dt.timedelta(days=RESERVE_STALE_DAYS)).strftime("%Y-%m-%d")
    fmap = files_map(db)
    items, total = [], 0
    for key, r in regs().items():
        if not can_edit(user, key):
            continue
        for year in years:
            rows = db.execute("SELECT id, seq, number, date, title, status FROM orders "
                              "WHERE type=? AND year=? AND deleted_at IS NULL ORDER BY date, seq, id",
                              (key, year)).fetchall()
            seqs = {x["seq"] for x in rows if x["seq"]}
            # обмеження — на випадок помилково вписаного величезного номера
            gaps = [n for n in range(1, min(max(seqs, default=0), 10000)) if n not in seqs]
            brief = lambda x: {"id": x["id"], "number": x["number"], "date": x["date"], "title": x["title"]}
            no_files = [brief(x) for x in rows if x["status"] != "reserved" and not fmap.get(x["id"])]
            reserves = [brief(x) for x in rows if x["status"] == "reserved" and x["date"] <= stale]
            if gaps or no_files or reserves:
                items.append({"key": key, "short": r["short"], "kind": r["kind"], "year": year,
                              "gaps": compress_ranges(gaps)[:100], "gap_count": len(gaps),
                              "no_files": no_files[:300], "no_files_count": len(no_files),
                              "reserves": reserves})
                total += len(gaps) + len(no_files) + len(reserves)
    # адміністратору: копія в хмару не вдалась або давно не робилась (лише якщо її налаштовано)
    cloud = cloud_status() if user["is_admin"] else None
    cloud_warn = None
    if cloud and (not cloud["ok"] or cloud["stale"]):
        cloud_warn = (f"Остання спроба {fmt_d(cloud['time'][:10])} {cloud['time'][11:16]} — помилка: {cloud['message']}"
                      if not cloud["ok"] else
                      f"Остання вдала копія {fmt_d(cloud['time'][:10])} {cloud['time'][11:16]} — понад {CLOUD_STALE_DAYS} дні тому")
        total += 1
    periods = check_periods(today)
    seen = set((user["checks_seen"] or "").split(","))
    return jsonify(items=items, total=total, years=years, stale_days=RESERVE_STALE_DAYS, cloud_warn=cloud_warn,
                   due=bool(total) and any(p not in seen for p in periods))


@app.post("/api/checks/seen")
def checks_seen():
    user = require_login()
    get_db().execute("UPDATE employees SET checks_seen=? WHERE id=?",
                     (",".join(check_periods()), user["id"]))
    return jsonify(ok=True)


# ---------------------------------------------------------------- Довідка (для всіх; розділи залежать від ролі)
def _read_help(name):
    with open(os.path.join(BASE, "help", name), encoding="utf-8") as f:
        return f.read()


@app.get("/api/help")
def help_data():
    import json
    user = current_user()
    adm = bool(user and user["is_admin"])
    sec = bool(user and (user["is_secretary"] or adm))
    hr = bool(user and (user["is_hr"] or adm))
    sections = [{"id": "all", "title": "Перегляд і пошук", "html": _read_help("all.html")}]
    if sec:
        sections.append({"id": "sec", "title": "Секретар", "html": _read_help("secretary.html") + _read_help("numbering.html")})
    if hr:
        sections.append({"id": "hr", "title": "Кадровик", "html": _read_help("hr.html") + _read_help("numbering.html")})
    if adm:
        sections.append({"id": "adm", "title": "Адміністратор", "html": _read_help("admin.html")})
    with open(os.path.join(BASE, "changelog.json"), encoding="utf-8") as f:
        log = json.load(f)
    return jsonify(version=log[0]["version"] if log else "", sections=sections,
                   changelog=log if adm else [])


# ---------------------------------------------------------------- Реєстри (адміністратор)
def registers_payload(db):
    used = {r["type"]: r["n"] for r in db.execute("SELECT type, COUNT(*) AS n FROM orders "
                                                    "GROUP BY type")}
    return [{**r, "used": used.get(r["key"], 0)} for r in regs().values()]


@app.get("/api/registers")
def registers_list():
    require_admin()
    return jsonify(registers=registers_payload(get_db()))


def read_register_form():
    d = request.get_json(silent=True) or {}
    return norm(d.get("short")), norm(d.get("title")), str(d.get("suffix") or "").strip()[:16], d


@app.post("/api/registers")
def registers_create():
    require_admin()
    db = get_db()
    short, title, suffix, d = read_register_form()
    kind = d.get("kind")
    if kind not in ("main", "hr"):
        return jsonify(error="Оберіть вид реєстру"), 400
    if not short or not title:
        return jsonify(error="Вкажіть коротку й повну назву реєстру"), 400
    nxt = (db.execute("SELECT COALESCE(MAX(sort), -1) + 1 FROM registers").fetchone()[0])
    key = "r" + uuid.uuid4().hex[:6]
    db.execute("INSERT INTO registers (key, kind, short, title, suffix, sort) VALUES (?,?,?,?,?,?)",
               (key, kind, short, title, suffix, nxt))
    g.pop("regs", None)
    audit(db, "reg_create", details=f"{short} ({'основний' if kind == 'main' else 'кадровий'}), суфікс «{suffix}»")
    return jsonify(key=key), 201


@app.put("/api/registers/<key>")
def registers_update(key):
    require_admin()
    db = get_db()
    if key not in regs():
        abort(404)
    short, title, suffix, _ = read_register_form()
    if not short or not title:
        return jsonify(error="Вкажіть коротку й повну назву реєстру"), 400
    old = regs()[key]
    db.execute("UPDATE registers SET short=?, title=?, suffix=? WHERE key=?",
               (short, title, suffix, key))
    g.pop("regs", None)
    ch = ([f"Вкладка: «{old['short']}» → «{short}»"] if short != old["short"] else []) + \
         (["Повну назву змінено"] if title != old["title"] else []) + \
         ([f"Суфікс: «{old['suffix']}» → «{suffix}»"] if suffix != old["suffix"] else [])
    if ch:
        audit(db, "reg_update", details=[short] + ch)
    return jsonify(ok=True)


@app.delete("/api/registers/<key>")
def registers_delete(key):
    require_admin()
    db = get_db()
    r = regs().get(key)
    if not r:
        abort(404)
    used = db.execute("SELECT COUNT(*) FROM orders WHERE type=?", (key,)).fetchone()[0]
    if used:
        return jsonify(error=f"У реєстрі є накази ({used}, разом із кошиком) — видалення неможливе"), 409
    if r["kind"] == "main" and sum(1 for x in regs().values() if x["kind"] == "main") <= 1:
        return jsonify(error="Має залишитись хоча б один основний реєстр"), 409
    db.execute("DELETE FROM registers WHERE key=?", (key,))
    g.pop("regs", None)
    audit(db, "reg_delete", details=r["short"])
    return jsonify(ok=True)


# ---------------------------------------------------------------- Експорт / резервна копія (адмін)
def status_text(row, lmaps):
    cancelled = any(x["kind"] == "cancels" for x in lmaps[1].get(row["id"], []))
    return STATUSES["void"] if cancelled else STATUSES.get(row["status"], row["status"])


def links_text(oid, lmaps):
    out, inc = lmaps
    parts = [f"{LINK_KINDS[x['kind']]} № {x['number']} від {fmt_d(x['date'])}" for x in out.get(oid, [])]
    parts += [f"{'Скасовано' if x['kind'] == 'cancels' else 'Змінено'} наказом № {x['number']} від {fmt_d(x['date'])}"
              for x in inc.get(oid, [])]
    return "; ".join(parts)


@app.get("/api/export.csv")
def export_csv():
    require_admin()
    db = get_db()
    cmap = creators_map(db)
    lmaps = links_maps(db)
    rows = db.execute(ORDER_SELECT + " WHERE o.deleted_at IS NULL ORDER BY o.date, COALESCE(o.seq,0), o.id").fetchall()
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["Номер", "Дата", "Реєстр", "Назва", "Хто створив", "Хто підписав", "Статус", "Зв'язки"])
    def safe(v):
        v = v if isinstance(v, str) else ("" if v is None else str(v))
        return "'" + v if v[:1] in ("=", "+", "-", "@") else v
    for r in rows:
        d = dt.datetime.strptime(r["date"], "%Y-%m-%d").strftime("%d.%m.%Y")
        w.writerow([safe(r["number"]), d, regs().get(r["type"], {}).get("short", r["type"]), safe(r["title"]),
                    safe("; ".join(c["name"] for c in cmap.get(r["id"], []))), safe(r["signer_name"] or ""),
                    status_text(r, lmaps), safe(links_text(r["id"], lmaps))])
    stamp = dt.date.today().strftime("%Y-%m-%d")
    return Response("﻿" + buf.getvalue(), mimetype="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f"attachment; filename=nakazy_{stamp}.csv"})


@app.get("/api/export-info")
def export_info():
    """Реєстри й роки, які може вивантажити секретар / кадровик / адміністратор."""
    user = require_login()
    if not user["is_admin"] and not (user["is_secretary"] or user["is_hr"]):
        abort(403)
    db = get_db()
    out = []
    for r in regs().values():
        if not can_view(user, r["key"]):
            continue
        years = [{"year": y["year"], "count": y["n"]} for y in db.execute(
            "SELECT year, COUNT(*) AS n FROM orders WHERE type=? AND deleted_at IS NULL "
            "GROUP BY year ORDER BY year DESC",
            (r["key"],))]
        out.append({"key": r["key"], "short": r["short"], "kind": r["kind"], "years": years})
    return jsonify(registers=out)


@app.get("/api/export.xlsx")
def export_xlsx():
    """Excel-файл наказів реєстру за рік (або по аркушу на кожен рік), без вкладень."""
    user = require_login()
    if not user["is_admin"] and not (user["is_secretary"] or user["is_hr"]):
        abort(403)
    type_ = request.args.get("type", "")
    if type_ not in regs() or not can_view(user, type_):
        abort(404)
    reg = regs()[type_]
    is_hr = reg["kind"] == "hr"
    year = request.args.get("year", "all")
    db = get_db()
    sql = ORDER_SELECT + " WHERE o.type=? AND o.deleted_at IS NULL"
    params = [type_]
    if year != "all":
        if not year.isdigit():
            abort(400)
        sql += " AND o.year=?"
        params.append(int(year))
    months = sorted({int(m) for m in request.args.get("months", "").split(",") if m.strip().isdigit() and 1 <= int(m) <= 12})
    if months and len(months) < 12:
        sql += f" AND CAST(substr(o.date, 6, 2) AS INTEGER) IN ({','.join('?' * len(months))})"
        params += months
    rows = db.execute(sql + " ORDER BY o.year, o.date, COALESCE(o.seq,0), o.id", params).fetchall()
    cmap = creators_map(db)
    lmaps = links_maps(db)
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return jsonify(error="Не встановлено openpyxl"), 500
    heads = ["Номер", "Дата", "Назва", "Хто створив", "Хто підписав", "Статус"] + ([] if is_hr else ["Зв'язки"])
    widths = [14, 12, 70, 28, 28, 14] + ([] if is_hr else [34])
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    by_year = {}
    for r in rows:
        by_year.setdefault(r["year"], []).append(r)
    if not by_year:
        by_year[int(year) if year.isdigit() else dt.date.today().year] = []
    for y in sorted(by_year):
        ws = wb.create_sheet(str(y))
        ws.append(heads)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E8ECF4")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r in by_year[y]:
            d = dt.datetime.strptime(r["date"], "%Y-%m-%d").date()
            line = [r["number"], d, r["title"], "; ".join(c["name"] for c in cmap.get(r["id"], [])),
                    r["signer_name"] or ""]
            line.append(status_text(r, lmaps))
            if not is_hr:
                line.append(links_text(r["id"], lmaps))
            ws.append(line)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            row[1].number_format = "DD.MM.YYYY"
            row[1].alignment = Alignment(horizontal="center", vertical="top")
            row[0].alignment = Alignment(horizontal="center", vertical="top")
            for c in row[2:]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            for c in row:
                if isinstance(c.value, str) and c.value[:1] in ("=", "+", "-", "@"):
                    c.data_type = "s"          # текст, а не формула
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{max(ws.max_row, 1)}"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    base = "nakazy_osnovni" if not is_hr else "nakazy_" + (re.sub(r"[^A-Za-z0-9.-]+", "_", reg["short"]).strip("_") or "kadrovi")
    mpart = "_mis_" + "-".join(f"{m:02d}" for m in months) if months and len(months) < 12 else ""
    name = f"{base}_{year if year != 'all' else 'usi_roky'}{mpart}.xlsx"
    return send_file(buf, as_attachment=True, download_name=name,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/api/backup")
def backup_db():
    """Консистентна копія бази (скани копіюйте з папки data/uploads)."""
    require_admin()
    tmp = os.path.join(DATA_DIR, f"backup_tmp_{uuid.uuid4().hex}.db")
    src = connect()
    dst = sqlite3.connect(tmp)
    src.backup(dst)
    dst.close()
    src.close()
    stamp = dt.date.today().strftime("%Y-%m-%d")
    resp = send_file(tmp, download_name=f"orders_{stamp}.db", as_attachment=True)
    resp.call_on_close(lambda: os.path.exists(tmp) and os.remove(tmp))
    return resp


# ---------------------------------------------------------------- Імпорт з Excel / CSV (адмін)
HEADER_SYNONYMS = {
    "number": ["номер", "№", "номер наказу", "№ наказу", "номер приказа", "№ приказа", "number"],
    "date": ["дата", "дата наказу", "дата приказа", "date"],
    "type": ["реєстр", "журнал", "тип", "вид", "вид наказу", "тип наказу", "вид приказа",
             "реестр", "register", "type"],
    "title": ["назва", "найменування", "зміст", "тема", "заголовок", "наименование", "название",
              "содержание", "тема наказу", "title"],
    "creators": ["хто створив", "створив", "виконавець", "відповідальний", "відповідальна особа",
                 "автор", "исполнитель", "ответственный", "кто создал", "создал", "responsible"],
    "signer": ["хто підписав", "підписав", "підписант", "кто подписал", "подписал", "подписант",
               "signer"],
    "status": ["статус", "стан", "status"],
}


def read_table(upload):
    name = (upload.filename or "").lower()
    raw = upload.read()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError:
            raise ValueError("Для імпорту .xlsx встановіть openpyxl (pip install openpyxl)")
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        return [list(r) for r in wb.active.iter_rows(values_only=True)]
    if name.endswith(".xls"):
        raise ValueError("Формат .xls не підтримується — збережіть файл як .xlsx або .csv")
    for enc in ("utf-8-sig", "cp1251", "cp866"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("Не вдалося визначити кодування CSV")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    except csv.Error:
        class dialect(csv.excel):   # копія, щоб не змінювати глобальний csv.excel
            delimiter = ";"
    return list(csv.reader(io.StringIO(text), dialect))


def guess_type(value, default):
    """Визначає реєстр за значенням колонки (ключ, коротка назва чи індекс реєстру, «основна…», «кадрові»)."""
    s = norm(value).casefold()
    if not s:
        return default
    for key, r in regs().items():
        short = r["short"].casefold()
        if s == key.casefold() or short in s or (len(s) >= 3 and s in short):
            return key
    if "основ" in s:
        return next((k for k, r in regs().items() if r["kind"] == "main"), default)
    if "кадр" in s or "особов" in s or "персонал" in s:
        if regs().get(default, {}).get("kind") == "hr":
            return default
        return next((k for k, r in regs().items() if r["kind"] == "hr"), default)
    return default


@app.post("/api/import")
def import_orders():
    user = require_admin()
    upload = request.files.get("file")
    if not upload:
        return jsonify(error="Оберіть файл"), 400
    default_type = request.form.get("default_type", "main")
    if default_type not in regs():
        default_type = next(iter(regs()), "main")
    try:
        table = read_table(upload)
    except ValueError as e:
        return jsonify(error=str(e)), 400
    table = [r for r in table if r and any(c not in (None, "") for c in r)]
    if len(table) < 2:
        return jsonify(error="У файлі немає даних"), 400
    headers = [str(c or "").replace("\xa0", " ").strip().lower() for c in table[0]]

    def header_keys(h):
        """Заголовок і його частини: «Назва/Зміст» -> {«назва/зміст», «назва», «зміст»}."""
        h = " ".join(h.split()).strip(" :.*")
        parts = [" ".join(x.split()).strip(" :.*") for x in re.split(r"[/,;|()\\]+", h)]
        return {h, *[x for x in parts if x]}

    col = {}
    for field, names in HEADER_SYNONYMS.items():
        for i, h in enumerate(headers):
            if field not in col and i not in col.values() and header_keys(h) & set(names):
                col[field] = i
    if "date" not in col or "title" not in col:
        return jsonify(error="Не знайдено обов'язкові колонки (Дата, Назва/Зміст). Знайдені "
                             "заголовки: " + ", ".join(h for h in headers if h)), 400

    def cell(row, field):
        i = col.get(field)
        return row[i] if i is not None and i < len(row) else None

    db = get_db()
    imported = skipped = 0
    problems = []
    db.execute("BEGIN IMMEDIATE")
    try:
        emp_cache = name_cache(db, "employees")
        sig_cache = name_cache(db, "signers")
        before_emp, before_sig = len(emp_cache), len(sig_cache)
        for n, row in enumerate(table[1:], start=2):
            date = parse_date(cell(row, "date"))
            title = str(cell(row, "title") or "").strip()
            reserved = "резерв" in str(cell(row, "status") or "").lower()
            if not date or (not title and not reserved):
                problems.append(f"Рядок {n}: немає коректної дати або назви")
                continue
            type_ = guess_type(cell(row, "type"), default_type)
            year = int(date[:4])
            number = str(cell(row, "number") or "").strip()
            if number.endswith(".0") and number[:-2].isdigit():
                number = number[:-2]
            if number:
                number = with_suffix(type_, number)   # 1 -> 1-к у кадровому реєстрі, як і при ручному створенні
                seq = seq_of(number)
            else:
                seq, number = next_number(db, type_, year)
            if duplicate_exists(db, type_, year, number):
                skipped += 1
                continue
            st_raw = str(cell(row, "status") or "").lower()
            status = "reserved" if reserved else \
                "void" if any(w in st_raw for w in ("втрат", "утрат", "скасов", "отмен", "недійс")) else "active"
            signer_name = norm(cell(row, "signer"))
            signer_id = ensure_name(db, "signers", sig_cache, signer_name) if signer_name else None
            ts = now_iso()
            cur = db.execute(
                "INSERT INTO orders (type, year, seq, number, date, title, responsible, status, "
                "signer_id, created_by, updated_by, created_at, updated_at) "
                "VALUES (?,?,?,?,?,?,'',?,?,?,?,?,?)",
                (type_, year, seq, number, date, title, status,
                 signer_id, user["id"], user["id"], ts, ts))
            for name in split_names(cell(row, "creators")):
                eid = ensure_name(db, "employees", emp_cache, name)
                db.execute("INSERT OR IGNORE INTO order_creators (order_id, employee_id) "
                           "VALUES (?,?)", (cur.lastrowid, eid))
            imported += 1
        audit(db, "import", details=f"Файл «{upload.filename}»: імпортовано {imported}, пропущено дублікатів {skipped}"
                                    + (f", рядків з помилками {len(problems)}" if problems else ""))
        db.execute("COMMIT")
    except Exception:
        db.execute("ROLLBACK")
        raise
    return jsonify(imported=imported, skipped=skipped, problems=problems[:20],
                   new_employees=len(emp_cache) - before_emp,
                   new_signers=len(sig_cache) - before_sig)


# файл розкладу копії для сервера — актуальний з моменту запуску програми
_conn = connect()
write_backup_schedule(_conn)
_conn.close()


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "8000"))
    print(f"Реєстр наказів: http://localhost:{port}  (у мережі: http://<IP-цього-комп'ютера>:{port})")
    app.run(host="0.0.0.0", port=port, threaded=True)
